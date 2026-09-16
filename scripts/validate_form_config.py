#!/usr/bin/env python3
import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter


SUPPORTED_TOKENS = {
    "project_name",
    "reimburser",
    "dingding_number",
    "reimbursement_type",
    "payee",
    "invoice_entity",
    "approval_company",
    "reimbursement_project",
    "contract_number",
    "reimbursement_reason",
    "business_line",
    "approved_amount",
    "total_amount",
    "expense_count",
    "date_range",
    "date_start",
    "date_end",
    "screenshot_print_pages",
    "invoice_print_pages",
    "expense_form_print_pages",
    "total_print_pages",
    "车费_amount",
    "车费_count",
    "餐杂费_amount",
    "餐杂费_count",
    "杂费_amount",
    "杂费_count",
    "交通费_amount",
    "交通费_count",
    "餐饮费_amount",
    "餐饮费_count",
    "设备费_amount",
    "设备费_count",
    "场地费_amount",
    "场地费_count",
    "演员费_amount",
    "演员费_count",
    "其他费用_amount",
    "其他费用_count",
}
TOKEN_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def parse_args():
    parser = argparse.ArgumentParser(description="Validate a reusable form_cells.json config before using it for reimbursement packaging.")
    parser.add_argument("--form-cells", required=True, help="Path to form_cells.json.")
    parser.add_argument("--template", help="Optional template override. Defaults to the template value in form_cells.json.")
    parser.add_argument("--out", help="Optional JSON report path. Defaults to stdout only.")
    return parser.parse_args()


def load_form_config(path):
    config_path = Path(path).expanduser().resolve()
    data = json.loads(config_path.read_text(encoding="utf-8"))
    if "reimbursement_form" in data:
        data = data.get("reimbursement_form") or {}
    if "recommended_manifest_reimbursement_form" in data:
        data = data.get("recommended_manifest_reimbursement_form") or {}
    return config_path, data


def resolve_template(value, base_dir):
    if not value:
        return None
    path = Path(value).expanduser()
    if not path.is_absolute():
        path = base_dir / path
    return path.resolve()


def writable_cell_ref(ws, cell_ref):
    for cell_range in ws.merged_cells.ranges:
        if cell_ref in cell_range:
            return f"{get_column_letter(cell_range.min_col)}{cell_range.min_row}"
    return cell_ref


def validate_cell_ref(ws, cell_ref):
    try:
        row, col = coordinate_to_tuple(cell_ref)
    except Exception as exc:
        return {"cell": cell_ref, "status": "invalid", "message": f"Invalid cell reference: {exc}"}
    if row < 1 or col < 1:
        return {"cell": cell_ref, "status": "invalid", "message": "Cell reference must be positive."}
    writable = writable_cell_ref(ws, cell_ref)
    status = "passed" if writable == cell_ref else "warning"
    message = "" if writable == cell_ref else f"Cell is inside a merged range; the actual writable top-left cell is {writable}."
    return {"cell": cell_ref, "writable_cell": writable, "status": status, "message": message}


def tokens_for(value):
    if not isinstance(value, str):
        return []
    return [match.group(1).strip() for match in TOKEN_RE.finditer(value)]


def validate_config(config_path, data, template_override=None):
    warnings = []
    errors = []
    template = resolve_template(template_override or data.get("template", ""), config_path.parent)
    if not template:
        errors.append("Missing template path. Add template to form_cells.json or pass --template.")
    elif not template.exists():
        errors.append(f"Template not found: {template}")
    elif template.suffix.lower() not in {".xlsx", ".xlsm"}:
        errors.append(f"Template must be .xlsx or .xlsm: {template}")

    cells = data.get("cells") or {}
    if not isinstance(cells, dict):
        errors.append("cells must be a JSON object mapping Excel cell references to values or tokens.")
        cells = {}
    if not cells:
        errors.append("cells is empty; no form fields will be filled.")
    category_rows = data.get("category_rows") or []
    if not isinstance(category_rows, list):
        errors.append("category_rows must be a JSON array of label_cell/amount_cell objects.")
        category_rows = []

    sheet = data.get("sheet", "")
    cell_checks = []
    category_row_checks = []
    unknown_tokens = []
    used_tokens = []
    sheet_names = []
    total_formula_detected = False
    if template and template.exists() and template.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(template, keep_vba=template.suffix.lower() == ".xlsm", data_only=False)
        sheet_names = wb.sheetnames
        if sheet and sheet not in wb.sheetnames:
            errors.append(f"Sheet not found: {sheet}. Available sheets: {', '.join(wb.sheetnames)}")
            ws = wb.active
        else:
            ws = wb[sheet] if sheet else wb.active
            if not sheet:
                warnings.append(f"No sheet configured; active sheet will be used: {ws.title}")
            total_formula_detected = any(
                cell.data_type == "f" and isinstance(cell.value, str) and "SUM(" in cell.value.upper()
                for row in ws.iter_rows()
                for cell in row
            )
        for cell_ref, value in cells.items():
            check = validate_cell_ref(ws, cell_ref)
            cell_checks.append({**check, "value": value})
            if check["status"] == "invalid":
                errors.append(f"{cell_ref}: {check['message']}")
            elif check["status"] == "warning":
                warnings.append(f"{cell_ref}: {check['message']}")
            for token in tokens_for(value):
                used_tokens.append(token)
                if token not in SUPPORTED_TOKENS and not token.endswith("_amount") and not token.endswith("_count"):
                    unknown_tokens.append(token)
        for index, row_config in enumerate(category_rows, start=1):
            if not isinstance(row_config, dict):
                errors.append(f"category_rows[{index}] must be an object.")
                continue
            row_check = {"row": index, "label_cell": row_config.get("label_cell", ""), "amount_cell": row_config.get("amount_cell", ""), "checks": []}
            for field in ("label_cell", "amount_cell"):
                cell_ref = row_config.get(field, "")
                if not cell_ref:
                    errors.append(f"category_rows[{index}].{field} is required.")
                    continue
                check = validate_cell_ref(ws, cell_ref)
                row_check["checks"].append({"field": field, **check})
                if check["status"] == "invalid":
                    errors.append(f"category_rows[{index}].{field}: {check['message']}")
                elif check["status"] == "warning":
                    warnings.append(f"category_rows[{index}].{field}: {check['message']}")
            category_row_checks.append(row_check)

    if unknown_tokens:
        errors.append(f"Unknown token(s): {', '.join(sorted(set(unknown_tokens)))}")
    if (
        "total_amount" not in used_tokens
        and not any(str(value).strip() == "{{total_amount}}" for value in cells.values())
        and not total_formula_detected
    ):
        warnings.append("No {{total_amount}} mapping found; confirm the template does not need a total amount field.")

    result = {
        "status": "failed" if errors else "passed_with_warnings" if warnings else "passed",
        "form_cells": str(config_path),
        "template": str(template) if template else "",
        "template_exists": bool(template and template.exists()),
        "sheet": sheet,
        "available_sheets": sheet_names,
        "cell_count": len(cells),
        "category_row_count": len(category_rows),
        "used_tokens": sorted(set(used_tokens)),
        "total_formula_detected": total_formula_detected,
        "cell_checks": cell_checks,
        "category_row_checks": category_row_checks,
        "warnings": warnings,
        "errors": errors,
        "next_step": "Fix errors before packaging. Review warnings against the visible template before team reuse." if errors or warnings else "Config is ready for package_from_folder.py --form-cells.",
    }
    return result


def main():
    args = parse_args()
    config_path, data = load_form_config(args.form_cells)
    result = validate_config(config_path, data, args.template)
    text = json.dumps(result, ensure_ascii=False, indent=2)
    if args.out:
        output = Path(args.out).expanduser().resolve()
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(text + "\n", encoding="utf-8")
    print(text)
    if result["status"] == "failed":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
