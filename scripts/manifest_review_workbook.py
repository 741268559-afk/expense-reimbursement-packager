#!/usr/bin/env python3
import argparse
import json
import re
import sys
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.worksheet.datavalidation import DataValidation


HEADERS = [
    "序号",
    "凭证编号",
    "日期",
    "时间",
    "用途",
    "费用类型",
    "金额",
    "凭证截图",
    "发票",
    "异常提示",
    "异常处理说明",
    "备注",
    "OCR文本",
]
EDITABLE_FIELDS = {
    "日期": "date",
    "时间": "time",
    "用途": "purpose",
    "费用类型": "expense_type",
    "金额": "amount",
    "异常处理说明": "exception_resolution",
    "备注": "notes",
}
DEFAULT_REVIEW_NOTES = {"", "NEEDS_REVIEW"}


def parse_args():
    parser = argparse.ArgumentParser(description="Export or apply an editable Excel review workbook for a reimbursement manifest.")
    subparsers = parser.add_subparsers(dest="command", required=True)

    export_parser = subparsers.add_parser("export", help="Create draft_manifest_review.xlsx from manifest JSON.")
    export_parser.add_argument("--manifest", required=True, help="Input manifest JSON.")
    export_parser.add_argument("--workbook", required=True, help="Output review workbook path.")

    apply_parser = subparsers.add_parser("apply", help="Apply review workbook edits back to manifest JSON.")
    apply_parser.add_argument("--manifest", required=True, help="Original manifest JSON.")
    apply_parser.add_argument("--workbook", required=True, help="Reviewed workbook path.")
    apply_parser.add_argument("--out", required=True, help="Output reviewed manifest JSON.")
    apply_parser.add_argument("--allow-new-categories", action="store_true", help="Allow expense_type values outside manifest categories.")
    return parser.parse_args()


def load_json(path):
    return json.loads(Path(path).expanduser().resolve().read_text(encoding="utf-8"))


def write_json(path, data):
    output = Path(path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return output


def stringify(value):
    if value is None:
        return ""
    return str(value).strip()


def normalize_date_value(value):
    if value in ("", None):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    text = stringify(value)
    match = re.search(r"(20\d{2})[-/.年](\d{1,2})[-/.月](\d{1,2})", text)
    if match:
        return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
    return text


def normalize_time_value(value):
    if value in ("", None):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%H:%M:%S")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    text = stringify(value)
    match = re.search(r"^(\d{1,2}):(\d{1,2})(?::(\d{1,2}))?$", text)
    if match:
        hour = int(match.group(1))
        minute = int(match.group(2))
        second = int(match.group(3) or 0)
        return f"{hour:02d}:{minute:02d}:{second:02d}"
    return text


def normalize_amount_value(value):
    if value in ("", None):
        return ""
    if isinstance(value, (int, float, Decimal)):
        amount = Decimal(str(value))
    else:
        text = stringify(value)
        text = text.replace(",", "").replace("¥", "").replace("￥", "").replace("元", "").replace("CNY", "").replace("RMB", "").strip()
        if not text:
            return ""
        try:
            amount = Decimal(text)
        except InvalidOperation:
            return stringify(value)
    return float(amount.quantize(Decimal("0.01")))


def invoice_names(entry):
    values = entry.get("invoices", []) or []
    if isinstance(values, str):
        values = [values]
    return "\n".join(Path(str(value)).name for value in values if str(value).strip())


def export_workbook(manifest_path, workbook_path):
    manifest = load_json(manifest_path)
    output = Path(workbook_path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "费用明细复核"
    ws.append(HEADERS)

    exceptions_by_voucher = {}
    for item in manifest.get("expense_exceptions", []) or []:
        exceptions_by_voucher.setdefault(item.get("voucher_id", ""), []).append(item.get("message", ""))

    for index, entry in enumerate(manifest.get("entries", []) or [], start=1):
        ws.append([
            index,
            entry.get("voucher_id", ""),
            entry.get("date", ""),
            entry.get("time", ""),
            entry.get("purpose", ""),
            entry.get("expense_type", ""),
            entry.get("amount", ""),
            Path(str(entry.get("screenshot", ""))).name,
            invoice_names(entry),
            "\n".join(exceptions_by_voucher.get(entry.get("voucher_id", ""), [])),
            entry.get("exception_resolution", ""),
            entry.get("notes", ""),
            entry.get("ocr_text", ""),
        ])

    header_fill = PatternFill("solid", fgColor="FFD966")
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    widths = {
        "A": 8,
        "B": 12,
        "C": 14,
        "D": 12,
        "E": 18,
        "F": 14,
        "G": 12,
        "H": 34,
        "I": 28,
        "J": 36,
        "K": 30,
        "L": 24,
        "M": 48,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{max(ws.max_row, 1)}"

    editable_columns = {HEADERS.index(header) + 1 for header in EDITABLE_FIELDS}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(HEADERS)):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if cell.column not in editable_columns:
                cell.fill = PatternFill("solid", fgColor="F2F2F2")
        ws.row_dimensions[cell.row].height = 54

    categories = [str(item).strip() for item in manifest.get("categories", []) or [] if str(item).strip()]
    if categories:
        cat_ws = wb.create_sheet("_categories")
        for row, value in enumerate(categories, start=1):
            cat_ws.cell(row=row, column=1, value=value)
        cat_ws.sheet_state = "hidden"
        formula = f"=_categories!$A$1:$A${len(categories)}"
        validation = DataValidation(type="list", formula1=formula, allow_blank=False)
        ws.add_data_validation(validation)
        category_column = ws.cell(row=1, column=HEADERS.index("费用类型") + 1).column_letter
        validation.add(f"{category_column}2:{category_column}{max(ws.max_row, 2)}")

    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=HEADERS.index("日期") + 1).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=HEADERS.index("时间") + 1).number_format = "hh:mm:ss"
        ws.cell(row=row, column=HEADERS.index("金额") + 1).number_format = '#,##0.00'

    wb.save(output)
    return {
        "status": "exported",
        "manifest": str(Path(manifest_path).expanduser().resolve()),
        "workbook": str(output),
        "rows": len(manifest.get("entries", []) or []),
    }


def apply_workbook(manifest_path, workbook_path, out_path, allow_new_categories=False):
    manifest = load_json(manifest_path)
    workbook = Path(workbook_path).expanduser().resolve()
    if not workbook.exists():
        raise SystemExit(f"Review workbook not found: {workbook}")
    wb = load_workbook(workbook, data_only=True)
    if "费用明细复核" not in wb.sheetnames:
        raise SystemExit("Review workbook must contain sheet: 费用明细复核")
    ws = wb["费用明细复核"]
    header_by_name = {}
    for column in range(1, ws.max_column + 1):
        value = stringify(ws.cell(row=1, column=column).value)
        if value:
            header_by_name[value] = column
    required_headers = ["序号", "日期", "时间", "用途", "费用类型", "金额"]
    missing_headers = [header for header in required_headers if header not in header_by_name]
    if missing_headers:
        raise SystemExit(f"Review workbook missing required headers: {', '.join(missing_headers)}")

    entries = manifest.get("entries", []) or []
    categories = {str(item).strip() for item in manifest.get("categories", []) or [] if str(item).strip()}
    errors = []
    applied = 0
    seen = set()

    for row in range(2, ws.max_row + 1):
        raw_index = ws.cell(row=row, column=header_by_name["序号"]).value
        if raw_index in ("", None):
            continue
        try:
            entry_index = int(raw_index)
        except (TypeError, ValueError):
            errors.append(f"第 {row} 行序号不是数字: {raw_index}")
            continue
        if entry_index < 1 or entry_index > len(entries):
            errors.append(f"第 {row} 行序号超出 manifest entries 范围: {entry_index}")
            continue
        if entry_index in seen:
            errors.append(f"第 {row} 行重复序号: {entry_index}")
            continue
        seen.add(entry_index)
        entry = entries[entry_index - 1]

        values = {}
        for header, field in EDITABLE_FIELDS.items():
            if header not in header_by_name:
                values[field] = entry.get(field, "")
                continue
            cell_value = ws.cell(row=row, column=header_by_name[header]).value
            if field == "date":
                values[field] = normalize_date_value(cell_value)
            elif field == "time":
                values[field] = normalize_time_value(cell_value)
            elif field == "amount":
                values[field] = normalize_amount_value(cell_value)
            else:
                values[field] = stringify(cell_value)

        row_errors = []
        if not values["date"]:
            row_errors.append("日期")
        if not values["purpose"]:
            row_errors.append("用途")
        if not values["expense_type"]:
            row_errors.append("费用类型")
        if values["amount"] == "":
            row_errors.append("金额")
        if not entry.get("screenshot"):
            row_errors.append("凭证截图")
        if row_errors:
            errors.append(f"第 {row} 行缺少: {', '.join(row_errors)}")
        if categories and values["expense_type"] and values["expense_type"] not in categories and not allow_new_categories:
            errors.append(f"第 {row} 行费用类型不在 manifest categories 中: {values['expense_type']}")

        entry.update({
            "date": values["date"],
            "time": values["time"],
            "purpose": values["purpose"],
            "expense_type": values["expense_type"],
            "amount": values["amount"],
            "exception_resolution": values["exception_resolution"],
        })
        complete_after_edit = all([
            entry.get("date"),
            entry.get("purpose"),
            entry.get("expense_type"),
            entry.get("amount") not in ("", None),
            entry.get("screenshot"),
        ])
        if complete_after_edit and values["notes"] in DEFAULT_REVIEW_NOTES:
            entry["notes"] = ""
        else:
            entry["notes"] = values["notes"]
        applied += 1

    if errors:
        print(json.dumps({
            "status": "failed",
            "manifest": str(Path(manifest_path).expanduser().resolve()),
            "workbook": str(workbook),
            "errors": errors,
        }, ensure_ascii=False, indent=2), file=sys.stderr)
        raise SystemExit(1)

    output = write_json(out_path, manifest)
    return {
        "status": "applied",
        "manifest": str(Path(manifest_path).expanduser().resolve()),
        "workbook": str(workbook),
        "out": str(output),
        "rows": applied,
    }


def main():
    args = parse_args()
    if args.command == "export":
        result = export_workbook(args.manifest, args.workbook)
    else:
        result = apply_workbook(args.manifest, args.workbook, args.out, args.allow_new_categories)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
