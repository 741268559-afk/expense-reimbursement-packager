#!/usr/bin/env python3
import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


def parse_args():
    parser = argparse.ArgumentParser(description="Run a deterministic self-test for the reimbursement packager skill.")
    parser.add_argument("--out", help="Optional output folder. Defaults to a temporary folder that is kept for inspection.")
    return parser.parse_args()


def run(cmd, cwd=None):
    result = subprocess.run(cmd, text=True, capture_output=True, cwd=cwd)
    if result.returncode != 0:
        if result.stdout:
            print(result.stdout, file=sys.stderr, end="")
        if result.stderr:
            print(result.stderr, file=sys.stderr, end="")
        raise SystemExit(result.returncode)
    return result.stdout


def make_screenshot(path, title, amount):
    image = Image.new("RGB", (900, 1400), "white")
    draw = ImageDraw.Draw(image)
    lines = [
        "Expense Screenshot",
        title,
        f"Amount: CNY {amount}",
        "Payment complete",
    ]
    y = 120
    for index, line in enumerate(lines):
        draw.text((90, y), line, fill="black")
        y += 90 if index == 0 else 70
    image.save(path)


def make_invoice_pdf(path):
    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    c.setFont("Helvetica-Bold", 18)
    c.drawString(72, height - 96, "Mock Invoice")
    c.setFont("Helvetica", 12)
    c.drawString(72, height - 132, "Invoice No: SELF-TEST-001")
    c.drawString(72, height - 156, "Amount: CNY 136.80")
    c.drawString(72, height - 180, "This file is generated for reimbursement packager self-test.")
    c.showPage()
    c.save()


def make_itinerary_pdf(path):
    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    c.setFont("Helvetica-Bold", 18)
    c.drawString(72, height - 96, "Mock Trip Table")
    c.setFont("Helvetica", 12)
    c.drawString(72, height - 132, "Supporting itinerary for invoice SELF-TEST-001")
    c.drawString(72, height - 156, "This is not an invoice and must not count toward coverage.")
    c.showPage()
    c.save()


def make_form_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "费用报销单"
    ws.merge_cells("A1:H1")
    ws["A1"] = "费用报销单"
    ws["A1"].font = Font(bold=True, size=18)
    ws["A1"].alignment = Alignment(horizontal="center")
    labels = [
        ("A3", "报销人", "B3"),
        ("D3", "项目", "E3"),
        ("A4", "发生日期", "B4"),
        ("D4", "单据张数", "E4"),
        ("A5", "报销金额", "B5"),
        ("A7", "交通费", "B7"),
        ("D7", "餐饮费", "E7"),
        ("A8", "设备费", "B8"),
    ]
    for label_cell, label, target_cell in labels:
        ws[label_cell] = label
        ws[label_cell].font = Font(bold=True)
        ws[target_cell] = ""
    for column in "ABCDEFGH":
        ws.column_dimensions[column].width = 16
    wb.save(path)


def write_manifest(path, screenshots, invoice, form_template, form_cells):
    approval_gap = path.parent / "钉钉信息缺口.md"
    approval_gap.write_text("# 钉钉信息缺口\n\n当前状态：完整\n", encoding="utf-8")
    manifest = {
        "workflow_version": 2,
        "project_name": "自检项目",
        "reimburser": "测试报销人",
        "categories": ["交通费", "餐饮费", "设备费", "场地费", "演员费", "其他费用"],
        "invoices": [str(invoice)],
        "invoice_requirement": {"policy": "full_amount"},
        "approval_requirement": {
            "policy": "required",
            "required_fields": ["dingding_number", "reimbursement_type", "payee", "invoice_entity", "approved_amount"],
        },
        "approval_metadata": {
            "dingding_number": "SELF-TEST-APPROVAL-001",
            "reimbursement_type": "项目报销",
            "reimburser": "测试报销人",
            "payee": "测试领款人",
            "invoice_entity": "测试开票对象",
            "approved_amount": 136.80,
        },
        "expense_confirmation": {
            "status": "confirmed_by_dingtalk",
            "confirmed_amount": "136.80",
            "dingding_number": "SELF-TEST-APPROVAL-001",
        },
        "workflow_reports": {"approval_gap": str(approval_gap)},
        "invoice_items": [{"path": str(invoice), "role": "original"}],
        "reimbursement_form": {
            "template": str(form_template),
            "sheet": form_cells.get("sheet", "费用报销单"),
            "output_name": "",
            "cells": form_cells.get("cells", {}),
        },
        "entries": [
            {
                "voucher_id": "FY001",
                "date": "2026-01-01",
                "time": "09:10:00",
                "purpose": "打车费",
                "expense_type": "交通费",
                "amount": 12.30,
                "screenshot": str(screenshots[0]),
                "invoices": [],
                "exception_resolution": "",
            },
            {
                "voucher_id": "FY002",
                "date": "2026-01-01",
                "time": "12:20:00",
                "purpose": "餐费",
                "expense_type": "餐饮费",
                "amount": 45.60,
                "screenshot": str(screenshots[1]),
                "invoices": [],
                "exception_resolution": "",
            },
            {
                "voucher_id": "FY003",
                "date": "2026-01-02",
                "time": "10:00:00",
                "purpose": "设备配送费",
                "expense_type": "设备费",
                "amount": 78.90,
                "screenshot": str(screenshots[2]),
                "invoices": [],
                "exception_resolution": "",
            },
        ],
    }
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main():
    args = parse_args()
    script_dir = Path(__file__).resolve().parent
    work_dir = Path(args.out).expanduser().resolve() if args.out else Path(tempfile.mkdtemp(prefix="expense_reimbursement_self_test_"))
    if args.out and work_dir.exists():
        shutil.rmtree(work_dir)
    work_dir.mkdir(parents=True, exist_ok=True)
    sources = work_dir / "sources"
    form_config = work_dir / "form_config"
    pack = work_dir / "pack"
    for folder in [sources, form_config, pack]:
        folder.mkdir(parents=True, exist_ok=True)

    screenshots = [
        sources / "001_滴滴_2026-01-01_12.30.png",
        sources / "002_餐费_2026-01-01_45.60.png",
        sources / "003_货拉拉_2026-01-02_78.90.png",
    ]
    make_screenshot(screenshots[0], "Taxi", "12.30")
    make_screenshot(screenshots[1], "Meal", "45.60")
    make_screenshot(screenshots[2], "Equipment delivery", "78.90")
    invoice = sources / "invoice_SELF_TEST_001.pdf"
    make_invoice_pdf(invoice)
    itinerary = sources / "行程报销单_SELF_TEST.pdf"
    make_itinerary_pdf(itinerary)
    form_template = form_config / "费用报销单_自检模板.xlsx"
    make_form_template(form_template)
    profile_path = work_dir / "reimbursement_profile.json"
    profile_path.write_text(json.dumps({
        "reimburser": "测试报销人",
        "approval_defaults": {
            "payee": "测试领款人",
            "invoice_entity": "测试开票对象",
        },
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    approval_path = work_dir / "approval_metadata.json"
    approval_path.write_text(json.dumps({
        "dingding_number": "SELF-TEST-APPROVAL-001",
        "reimbursement_type": "项目报销",
        "reimburser": "测试报销人",
        "payee": "测试领款人",
        "invoice_entity": "测试开票对象",
        "approved_amount": 136.80,
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    form_inspection = form_config / "form_inspection.json"
    form_cells_path = form_config / "form_cells.json"
    run([
        sys.executable,
        str(script_dir / "inspect_reimbursement_form.py"),
        "--template",
        str(form_template),
        "--out",
        str(form_inspection),
        "--cells-out",
        str(form_cells_path),
        "--relative-template",
    ])
    form_cells = json.loads(form_cells_path.read_text(encoding="utf-8"))
    if form_cells.get("template") != form_template.name:
        raise SystemExit(f"Expected relative template path in form_cells.json, got: {form_cells.get('template')}")
    validation_path = form_config / "form_config_validation.json"
    validation_text = run([
        sys.executable,
        str(script_dir / "validate_form_config.py"),
        "--form-cells",
        str(form_cells_path),
        "--out",
        str(validation_path),
    ])
    validation = json.loads(validation_text)
    if validation.get("status") not in {"passed", "passed_with_warnings"} or validation.get("cell_count") != len(form_cells.get("cells", {})):
        raise SystemExit(f"Unexpected form config validation: {validation}")
    dynamic_form_cells_path = form_config / "form_cells_dynamic.json"
    dynamic_form_cells = {
        **form_cells,
        "cells": {
            token: cell_ref
            for token, cell_ref in form_cells.get("cells", {}).items()
            if token not in {"交通费_amount", "餐饮费_amount", "设备费_amount"}
        },
        "category_rows": [
            {"label_cell": "A7", "amount_cell": "B7"},
            {"label_cell": "D7", "amount_cell": "E7"},
            {"label_cell": "A8", "amount_cell": "B8"},
        ],
    }
    dynamic_form_cells_path.write_text(
        json.dumps(dynamic_form_cells, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    dynamic_validation_path = form_config / "form_config_dynamic_validation.json"
    dynamic_validation_text = run([
        sys.executable,
        str(script_dir / "validate_form_config.py"),
        "--form-cells",
        str(dynamic_form_cells_path),
        "--out",
        str(dynamic_validation_path),
    ])
    dynamic_validation = json.loads(dynamic_validation_text)
    if dynamic_validation.get("status") not in {"passed", "passed_with_warnings"} or dynamic_validation.get("category_row_count") != 3:
        raise SystemExit(f"Unexpected dynamic form config validation: {dynamic_validation}")
    readiness_without_form_text = run([
        sys.executable,
        str(script_dir / "check_readiness.py"),
        "--out",
        str(work_dir / "readiness_without_form.json"),
    ])
    readiness_without_form = json.loads(readiness_without_form_text)
    if readiness_without_form.get("status") != "needs_form_template":
        raise SystemExit(f"Unexpected readiness without form status: {readiness_without_form}")
    readiness_with_form_text = run([
        sys.executable,
        str(script_dir / "check_readiness.py"),
        "--form-cells",
        str(form_cells_path),
        "--out",
        str(work_dir / "readiness_with_form.json"),
    ])
    readiness_with_form = json.loads(readiness_with_form_text)
    if readiness_with_form.get("status") != "ready":
        raise SystemExit(f"Unexpected readiness with form status: {readiness_with_form}")
    onboard_config = work_dir / "onboard_config"
    onboarding_text = run([
        sys.executable,
        str(script_dir / "onboard_reimbursement_form.py"),
        "--template",
        str(form_template),
        "--out",
        str(onboard_config),
        "--copy-template",
    ])
    onboarding = json.loads(onboarding_text)
    if onboarding.get("status") != "ready_for_review" or onboarding.get("readiness_status") != "ready":
        raise SystemExit(f"Unexpected onboarding result: {onboarding}")
    if not (onboard_config / "readiness_report.json").exists():
        raise SystemExit("Expected onboard_reimbursement_form.py to write readiness_report.json.")
    team_like_root = work_dir / "team_like_root"
    team_like_config = team_like_root / "02_form_template_config"
    team_like_config.mkdir(parents=True, exist_ok=True)
    shutil.copy2(form_template, team_like_config / form_template.name)
    shutil.copy2(dynamic_form_cells_path, team_like_config / "form_cells.json")
    auto_config_pack = work_dir / "auto_config_pack"
    run([
        sys.executable,
        str(script_dir / "package_from_folder.py"),
        "--input",
        str(sources),
        "--project-name",
        "自检项目",
        "--profile",
        str(profile_path),
        "--approval-metadata",
        str(approval_path),
        "--out",
        str(auto_config_pack),
        "--ocr",
        "none",
        "--review-policy",
        "never",
    ], cwd=team_like_root)
    auto_package_result = json.loads((auto_config_pack / "package_result.json").read_text(encoding="utf-8"))
    if auto_package_result.get("status") != "passed" or not auto_package_result.get("expense_form"):
        raise SystemExit(f"Unexpected auto config package_result: {auto_package_result}")
    if auto_package_result.get("form_cells") != str((team_like_config / "form_cells.json").resolve()):
        raise SystemExit(f"Expected auto-discovered form_cells in package_result: {auto_package_result.get('form_cells')}")
    if auto_package_result.get("profile_source") != str(profile_path.resolve()):
        raise SystemExit(f"Expected reusable profile source in package_result: {auto_package_result.get('profile_source')}")
    dynamic_form_workbook = load_workbook(auto_package_result["expense_form"], data_only=False)
    dynamic_form_sheet = dynamic_form_workbook["费用报销单"]
    expected_dynamic_rows = [
        ("A7", "B7", "交通费", 12.30),
        ("D7", "E7", "餐饮费", 45.60),
        ("A8", "B8", "设备费", 78.90),
    ]
    for label_cell, amount_cell, expected_label, expected_amount in expected_dynamic_rows:
        if dynamic_form_sheet[label_cell].value != expected_label or abs(float(dynamic_form_sheet[amount_cell].value) - expected_amount) > 0.001:
            raise SystemExit(
                f"Dynamic category row mismatch at {label_cell}/{amount_cell}: "
                f"{dynamic_form_sheet[label_cell].value!r}, {dynamic_form_sheet[amount_cell].value!r}"
            )

    approval_gate_pack = work_dir / "approval_gate_pack"
    run([
        sys.executable,
        str(script_dir / "package_from_folder.py"),
        "--input",
        str(sources),
        "--project-name",
        "自检项目",
        "--reimburser",
        "测试报销人",
        "--out",
        str(approval_gate_pack),
        "--form-cells",
        str(form_cells_path),
        "--ocr",
        "none",
        "--review-policy",
        "never",
    ])
    approval_gate_result = json.loads((approval_gate_pack / "package_result.json").read_text(encoding="utf-8"))
    if (
        approval_gate_result.get("status") != "needs_approval"
        or not Path(approval_gate_result.get("confirmation_report", "")).exists()
        or not Path(approval_gate_result.get("approval_gap_report", "")).exists()
        or approval_gate_result.get("finance_zip")
    ):
        raise SystemExit(f"Unexpected DingTalk approval gate result: {approval_gate_result}")

    exception_manifest = work_dir / "exception_manifest.json"
    write_manifest(exception_manifest, screenshots, invoice, form_template, form_cells)
    exception_data = json.loads(exception_manifest.read_text(encoding="utf-8"))
    exception_data["entries"][0]["ocr_text"] = "退款成功 CNY 12.30"
    exception_manifest.write_text(json.dumps(exception_data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    exception_result = json.loads(run([
        sys.executable,
        str(script_dir / "prepare_expense_confirmation.py"),
        "--manifest",
        str(exception_manifest),
        "--out-dir",
        str(work_dir / "exception_reports"),
        "--update-manifest",
    ]))
    if exception_result.get("status") != "needs_review" or exception_result.get("unresolved_exception_count") != 1:
        raise SystemExit(f"Expected refund exception to block packaging: {exception_result}")
    exception_data = json.loads(exception_manifest.read_text(encoding="utf-8"))
    exception_data["entries"][0]["exception_resolution"] = "已确认该截图为未退款的测试凭证。"
    exception_manifest.write_text(json.dumps(exception_data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    resolved_exception_result = json.loads(run([
        sys.executable,
        str(script_dir / "prepare_expense_confirmation.py"),
        "--manifest",
        str(exception_manifest),
        "--out-dir",
        str(work_dir / "exception_reports"),
        "--update-manifest",
    ]))
    if resolved_exception_result.get("status") != "ready" or resolved_exception_result.get("unresolved_exception_count") != 0:
        raise SystemExit(f"Expected explained refund exception to clear the gate: {resolved_exception_result}")

    # Exercise package_from_folder enough to prove reusable relative form config resolves.
    draft_pack = work_dir / "draft_pack"
    run([
        sys.executable,
        str(script_dir / "package_from_folder.py"),
        "--input",
        str(sources),
        "--project-name",
        "自检项目",
        "--reimburser",
        "测试报销人",
        "--profile",
        str(profile_path),
        "--approval-metadata",
        str(approval_path),
        "--out",
        str(draft_pack),
        "--form-cells",
        str(form_cells_path),
        "--ocr",
        "none",
        "--review-policy",
        "always",
    ])
    preflight_report = draft_pack / "preflight_report.json"
    if not preflight_report.exists():
        raise SystemExit("Expected package_from_folder.py to write preflight_report.json.")
    preflight = json.loads(preflight_report.read_text(encoding="utf-8"))
    counts = preflight.get("counts", {})
    if counts.get("expense_screenshots") != 3 or counts.get("invoices") != 1 or counts.get("supporting_documents") != 1:
        raise SystemExit(f"Unexpected preflight counts: {counts}")
    review_result_path = draft_pack / "package_result.json"
    if not review_result_path.exists():
        raise SystemExit("Expected review run to write package_result.json.")
    review_result = json.loads(review_result_path.read_text(encoding="utf-8"))
    if review_result.get("status") != "needs_review":
        raise SystemExit(f"Expected review package_result status needs_review, got: {review_result.get('status')}")
    if "package_reviewed_workbook.py" not in review_result.get("next_step", ""):
        raise SystemExit(f"Expected review next_step to mention package_reviewed_workbook.py, got: {review_result.get('next_step')}")
    review_workbook = draft_pack / "draft_manifest_review.xlsx"
    if not review_workbook.exists() or Path(review_result.get("manifest_review_workbook", "")).resolve() != review_workbook.resolve():
        raise SystemExit("Expected review run to write and index draft_manifest_review.xlsx.")
    reviewed_manifest = draft_pack / "reviewed_manifest_from_workbook.json"
    review_apply_text = run([
        sys.executable,
        str(script_dir / "manifest_review_workbook.py"),
        "apply",
        "--manifest",
        str(draft_pack / "draft_manifest.json"),
        "--workbook",
        str(review_workbook),
        "--out",
        str(reviewed_manifest),
    ])
    review_apply = json.loads(review_apply_text)
    if review_apply.get("status") != "applied" or review_apply.get("rows") != 3:
        raise SystemExit(f"Unexpected manifest review apply result: {review_apply}")
    reviewed_data = json.loads(reviewed_manifest.read_text(encoding="utf-8"))
    if any(entry.get("notes") for entry in reviewed_data.get("entries", [])):
        raise SystemExit("Expected reviewed manifest from workbook to clear review notes for complete rows.")
    reviewed_pack = work_dir / "reviewed_workbook_pack"
    run([
        sys.executable,
        str(script_dir / "package_reviewed_workbook.py"),
        "--manifest",
        str(draft_pack / "draft_manifest.json"),
        "--workbook",
        str(review_workbook),
        "--out",
        str(reviewed_pack),
        "--preflight-report",
        str(preflight_report),
        "--form-inspection",
        str(draft_pack / "form_inspection.json"),
    ])
    reviewed_package_result_path = reviewed_pack / "package_result.json"
    if not reviewed_package_result_path.exists():
        raise SystemExit("Expected package_reviewed_workbook.py to write package_result.json.")
    reviewed_package_result = json.loads(reviewed_package_result_path.read_text(encoding="utf-8"))
    if reviewed_package_result.get("status") != "passed" or reviewed_package_result.get("verification_status") != "passed":
        raise SystemExit(f"Unexpected reviewed workbook package_result: {reviewed_package_result}")
    if Path(reviewed_package_result.get("manifest_review_workbook", "")).resolve() != review_workbook.resolve():
        raise SystemExit("Expected reviewed workbook package_result to keep the manifest_review_workbook path.")

    one_command_pack = work_dir / "one_command_pack"
    run([
        sys.executable,
        str(script_dir / "package_from_folder.py"),
        "--input",
        str(sources),
        "--project-name",
        "自检项目",
        "--reimburser",
        "测试报销人",
        "--approval-metadata",
        str(approval_path),
        "--out",
        str(one_command_pack),
        "--form-cells",
        str(form_cells_path),
        "--ocr",
        "none",
        "--review-policy",
        "never",
    ])
    package_result_path = one_command_pack / "package_result.json"
    if not package_result_path.exists():
        raise SystemExit("Expected one-command run to write package_result.json.")
    package_result = json.loads(package_result_path.read_text(encoding="utf-8"))
    if package_result.get("status") != "passed" or package_result.get("verification_status") != "passed":
        raise SystemExit(f"Unexpected one-command package_result: {package_result}")
    if package_result.get("supporting_document_count") != 1:
        raise SystemExit(f"Expected one supporting itinerary document in the final pack: {package_result}")
    for key in [
        "delivery_checklist",
        "finance_zip",
        "one_click_print_pdf",
        "reimbursement_workbook",
        "expense_form",
        "confirmation_report",
        "exception_report",
        "approval_gap_report",
        "audit_summary",
    ]:
        if not package_result.get(key):
            raise SystemExit(f"Missing {key} in package_result.json")
    audit_text = Path(package_result["audit_summary"]).read_text(encoding="utf-8")
    if "自动校验：全部通过" not in audit_text:
        raise SystemExit(f"Expected finalized audit summary status, got: {audit_text}")

    gap_sources = work_dir / "gap_sources"
    gap_sources.mkdir(parents=True, exist_ok=True)
    for screenshot in screenshots:
        shutil.copy2(screenshot, gap_sources / screenshot.name)

    gap_pack = work_dir / "gap_pack"
    run([
        sys.executable,
        str(script_dir / "package_from_folder.py"),
        "--input",
        str(gap_sources),
        "--project-name",
        "自检项目",
        "--reimburser",
        "测试报销人",
        "--approval-metadata",
        str(approval_path),
        "--out",
        str(gap_pack),
        "--form-cells",
        str(form_cells_path),
        "--ocr",
        "none",
        "--review-policy",
        "never",
    ])
    gap_result = json.loads((gap_pack / "package_result.json").read_text(encoding="utf-8"))
    if (
        gap_result.get("status") != "needs_invoices"
        or gap_result.get("invoice_required_amount") != "136.80"
        or gap_result.get("invoice_recognized_amount") != "0.00"
        or gap_result.get("invoice_missing_amount") != "136.80"
        or not Path(gap_result.get("invoice_gap_report", "")).exists()
    ):
        raise SystemExit(f"Unexpected first-stage invoice gap result: {gap_result}")
    if gap_result.get("finance_zip") or gap_result.get("one_click_print_pdf"):
        raise SystemExit("Invoice-gap stage must not generate a final finance zip or one-click print PDF.")

    replacement_pack = work_dir / "replacement_invoice_pack"
    run([
        sys.executable,
        str(script_dir / "package_from_folder.py"),
        "--input",
        str(gap_sources),
        "--replacement-invoice-input",
        str(invoice),
        "--project-name",
        "自检项目",
        "--reimburser",
        "测试报销人",
        "--approval-metadata",
        str(approval_path),
        "--out",
        str(replacement_pack),
        "--form-cells",
        str(form_cells_path),
        "--ocr",
        "none",
        "--review-policy",
        "never",
    ])
    replacement_result = json.loads((replacement_pack / "package_result.json").read_text(encoding="utf-8"))
    if (
        replacement_result.get("status") != "passed"
        or replacement_result.get("verification_status") != "passed"
        or replacement_result.get("invoice_coverage_status") != "covered"
        or replacement_result.get("invoice_missing_amount") != "0.00"
        or replacement_result.get("replacement_invoice_count") != 1
        or replacement_result.get("replacement_invoice_amount") != "136.80"
        or not Path(replacement_result.get("invoice_print_pdf", "")).exists()
        or not Path(replacement_result.get("one_click_print_pdf", "")).exists()
    ):
        raise SystemExit(f"Unexpected replacement-invoice package result: {replacement_result}")
    replacement_verification = json.loads((replacement_pack / "verification.json").read_text(encoding="utf-8"))
    required_replacement_checks = {
        "finance zip has invoice coverage JSON",
        "finance zip has invoice gap report",
        "finance zip replacement invoice count",
    }
    passed_replacement_checks = {
        item.get("name")
        for item in replacement_verification.get("checks", [])
        if item.get("status") == "passed"
    }
    if replacement_verification.get("status") != "passed" or not required_replacement_checks.issubset(passed_replacement_checks):
        raise SystemExit(f"Replacement-invoice verification did not prove final package contents: {replacement_verification}")

    duplicate_invoice = work_dir / "invoice_SELF_TEST_001_copy.pdf"
    shutil.copy2(invoice, duplicate_invoice)
    duplicate_manifest = work_dir / "duplicate_invoice_manifest.json"
    write_manifest(duplicate_manifest, screenshots, invoice, form_template, form_cells)
    duplicate_manifest_data = json.loads(duplicate_manifest.read_text(encoding="utf-8"))
    duplicate_manifest_data["invoices"] = [str(invoice), str(duplicate_invoice)]
    duplicate_manifest_data["invoice_items"] = [
        {"path": str(invoice), "role": "original"},
        {"path": str(duplicate_invoice), "role": "replacement"},
    ]
    duplicate_manifest.write_text(json.dumps(duplicate_manifest_data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    duplicate_coverage_path = work_dir / "duplicate_invoice_coverage.json"
    run([
        sys.executable,
        str(script_dir / "analyze_invoice_coverage.py"),
        "--manifest",
        str(duplicate_manifest),
        "--out",
        str(duplicate_coverage_path),
        "--update-manifest",
    ])
    duplicate_coverage = json.loads(duplicate_coverage_path.read_text(encoding="utf-8"))
    duplicate_items = [item for item in duplicate_coverage.get("invoice_items", []) if item.get("status") == "duplicate"]
    if (
        duplicate_coverage.get("status") != "needs_invoice_review"
        or duplicate_coverage.get("recognized_amount") != "136.80"
        or len(duplicate_items) != 1
    ):
        raise SystemExit(f"Duplicate invoice content was not rejected correctly: {duplicate_coverage}")

    manifest = work_dir / "manifest.json"
    write_manifest(manifest, screenshots, invoice, form_template, form_cells)
    run([
        sys.executable,
        str(script_dir / "prepare_expense_confirmation.py"),
        "--manifest",
        str(manifest),
        "--out-dir",
        str(work_dir),
        "--update-manifest",
    ])
    run([
        sys.executable,
        str(script_dir / "analyze_invoice_coverage.py"),
        "--manifest",
        str(manifest),
        "--out",
        str(work_dir / "manual_invoice_coverage.json"),
        "--update-manifest",
    ])
    run([
        sys.executable,
        str(script_dir / "build_reimbursement_pack.py"),
        "--manifest",
        str(manifest),
        "--out",
        str(pack),
    ])
    verification_text = run([
        sys.executable,
        str(script_dir / "verify_reimbursement_pack.py"),
        "--pack",
        str(pack),
        "--manifest",
        str(manifest),
    ])
    verification = json.loads(verification_text)
    if verification.get("status") != "passed":
        raise SystemExit("Self-test verification failed.")

    reimbursement_workbook = next((pack / "报销表").glob("*报销明细表.xlsx"))
    reimbursement_backup = work_dir / "reimbursement_before_tamper.xlsx"
    shutil.copy2(reimbursement_workbook, reimbursement_backup)
    tampered_workbook = load_workbook(reimbursement_workbook)
    tampered_sheet = tampered_workbook["报销明细"]
    header_row = next(
        row
        for row in range(1, tampered_sheet.max_row + 1)
        if tampered_sheet.cell(row, 1).value in {"序号", "凭证编号"}
    )
    first_amount_cell = tampered_sheet.cell(header_row + 1, 6)
    first_amount_cell.value = float(first_amount_cell.value) + 0.01
    tampered_workbook.save(reimbursement_workbook)
    tamper_result = subprocess.run(
        [
            sys.executable,
            str(script_dir / "verify_reimbursement_pack.py"),
            "--pack",
            str(pack),
            "--manifest",
            str(manifest),
        ],
        text=True,
        capture_output=True,
    )
    tamper_verification = json.loads((pack / "verification.json").read_text(encoding="utf-8"))
    tamper_detected = any(
        item.get("name") == "reimbursement workbook rows match manifest exactly"
        and item.get("status") == "failed"
        for item in tamper_verification.get("checks", [])
    )
    if tamper_result.returncode == 0 or tamper_verification.get("status") != "failed" or not tamper_detected:
        raise SystemExit("Self-test expected strict financial reconciliation to reject a tampered amount.")
    shutil.copy2(reimbursement_backup, reimbursement_workbook)
    reimbursement_backup.unlink()
    verification_text = run([
        sys.executable,
        str(script_dir / "verify_reimbursement_pack.py"),
        "--pack",
        str(pack),
        "--manifest",
        str(manifest),
    ])
    verification = json.loads(verification_text)
    if verification.get("status") != "passed":
        raise SystemExit("Self-test failed to recover after the tamper-detection check.")

    print(json.dumps({
        "status": "passed",
        "work_dir": str(work_dir),
        "pack": str(pack),
        "preflight_report": str(preflight_report),
        "package_result": str(package_result_path),
        "form_config_validation": str(validation_path),
        "dynamic_form_config_validation": str(dynamic_validation_path),
        "readiness_without_form": str(work_dir / "readiness_without_form.json"),
        "readiness_with_form": str(work_dir / "readiness_with_form.json"),
        "onboarding_report": str(onboard_config / "onboarding_report.json"),
        "onboarding_readiness_report": str(onboard_config / "readiness_report.json"),
        "auto_config_package_result": str(auto_config_pack / "package_result.json"),
        "approval_gate_package_result": str(approval_gate_pack / "package_result.json"),
        "expense_exception_report": str(work_dir / "exception_reports" / "报销异常清单.md"),
        "invoice_gap_package_result": str(gap_pack / "package_result.json"),
        "replacement_invoice_package_result": str(replacement_pack / "package_result.json"),
        "duplicate_invoice_coverage": str(duplicate_coverage_path),
        "verification": str(pack / "verification.json"),
        "checks": len(verification.get("checks", [])),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
