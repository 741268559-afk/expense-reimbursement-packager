#!/usr/bin/env python3
import argparse
import json
import os
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


FIELD_RULES = [
    ("reimburser", "{{reimburser}}", ("报销人", "申请人", "经办人", "姓名"), 0.92),
    ("project_name", "{{project_name}}", ("项目", "事由", "用途", "摘要", "说明"), 0.82),
    ("date_range", "{{date_range}}", ("日期范围", "报销日期", "起止日期", "发生日期"), 0.86),
    ("total_amount", "{{total_amount}}", ("报销金额", "合计金额", "金额合计", "总金额", "总计", "费用合计", "报销总额", "金额"), 0.95),
    ("expense_count", "{{expense_count}}", ("单据张数", "附件张数", "票据张数", "张数"), 0.88),
    ("交通费_amount", "{{交通费_amount}}", ("交通费", "车费", "交通"), 0.78),
    ("餐饮费_amount", "{{餐饮费_amount}}", ("餐饮费", "餐费", "餐饮", "餐杂费"), 0.78),
    ("设备费_amount", "{{设备费_amount}}", ("设备费", "设备", "器材"), 0.76),
    ("场地费_amount", "{{场地费_amount}}", ("场地费", "场租", "场地"), 0.76),
    ("演员费_amount", "{{演员费_amount}}", ("演员费", "演员", "模特"), 0.76),
    ("其他费用_amount", "{{其他费用_amount}}", ("其他费用", "杂费", "其他"), 0.72),
]


def parse_args():
    parser = argparse.ArgumentParser(description="Inspect an Excel reimbursement-form template and suggest manifest cell mappings.")
    parser.add_argument("--template", required=True, help="Path to the 费用报销单 Excel template.")
    parser.add_argument("--sheet", help="Optional sheet name to inspect. Defaults to all sheets.")
    parser.add_argument("--out", help="Optional JSON output path. Defaults to stdout.")
    parser.add_argument("--cells-out", help="Optional reusable form-cells JSON output for package_from_folder.py --form-cells.")
    parser.add_argument("--relative-template", action="store_true", help="Store the template path in --cells-out relative to the JSON file location.")
    parser.add_argument("--max-cells", type=int, default=300, help="Maximum non-empty cells to include per sheet.")
    return parser.parse_args()


def clean_text(value):
    if value is None:
        return ""
    text = str(value).strip()
    return re.sub(r"\s+", " ", text)


def merged_range_for(ws, row, col):
    coord = f"{get_column_letter(col)}{row}"
    for cell_range in ws.merged_cells.ranges:
        if coord in cell_range:
            return cell_range
    return None


def writable_ref(ws, row, col):
    cell_range = merged_range_for(ws, row, col)
    if cell_range:
        return f"{get_column_letter(cell_range.min_col)}{cell_range.min_row}"
    return f"{get_column_letter(col)}{row}"


def is_blank(ws, row, col):
    ref = writable_ref(ws, row, col)
    return clean_text(ws[ref].value) == ""


def label_bounds(ws, row, col):
    cell_range = merged_range_for(ws, row, col)
    if cell_range:
        return cell_range.min_row, cell_range.max_row, cell_range.min_col, cell_range.max_col
    return row, row, col, col


def find_target_cell(ws, row, col):
    min_row, max_row, min_col, max_col = label_bounds(ws, row, col)
    seen = set()
    candidates = []
    for offset in range(1, 6):
        candidates.append((min_row, max_col + offset))
    for offset in range(1, 4):
        candidates.append((max_row + offset, min_col))
    for target_row, target_col in candidates:
        if target_row < 1 or target_col < 1 or target_row > ws.max_row + 2 or target_col > ws.max_column + 2:
            continue
        ref = writable_ref(ws, target_row, target_col)
        if ref in seen:
            continue
        seen.add(ref)
        if is_blank(ws, target_row, target_col):
            return ref
    return ""


def non_empty_cells(ws, max_cells):
    cells = []
    for row in ws.iter_rows():
        for cell in row:
            value = clean_text(cell.value)
            if value:
                cells.append({"cell": cell.coordinate, "value": value})
                if len(cells) >= max_cells:
                    return cells
    return cells


def score_rule(text, keywords, base_score):
    normalized = text.replace(" ", "")
    for keyword in keywords:
        if keyword == "杂费" and "餐杂费" in normalized and normalized != "杂费":
            continue
        if keyword in normalized:
            return base_score
    return 0


def mapping_candidates(ws):
    candidates = []
    for row in ws.iter_rows():
        for cell in row:
            text = clean_text(cell.value)
            if not text:
                continue
            for field, token, keywords, base_score in FIELD_RULES:
                score = score_rule(text, keywords, base_score)
                if not score:
                    continue
                target = find_target_cell(ws, cell.row, cell.column)
                candidates.append({
                    "field": field,
                    "value": token,
                    "label_cell": cell.coordinate,
                    "label_text": text,
                    "target_cell": target,
                    "confidence": score if target else round(score - 0.2, 2),
                })
    candidates.sort(key=lambda item: (-item["confidence"], item["field"], item["label_cell"]))
    return candidates


def recommended_cells(candidates):
    cells = {}
    used_fields = set()
    used_targets = set()
    for candidate in candidates:
        target = candidate.get("target_cell")
        field = candidate.get("field")
        if not target or field in used_fields or target in used_targets:
            continue
        cells[target] = candidate["value"]
        used_fields.add(field)
        used_targets.add(target)
    return cells


def inspect_workbook(path, sheet_name, max_cells):
    wb = load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm", data_only=False)
    sheets = [wb[sheet_name]] if sheet_name else list(wb.worksheets)
    inspected = []
    recommended_sheet = ""
    recommended_mapping = {}
    for ws in sheets:
        candidates = mapping_candidates(ws)
        sheet_mapping = recommended_cells(candidates)
        if sheet_mapping and not recommended_mapping:
            recommended_sheet = ws.title
            recommended_mapping = sheet_mapping
        inspected.append({
            "title": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "merged_ranges": [str(item) for item in ws.merged_cells.ranges],
            "non_empty_cells": non_empty_cells(ws, max_cells),
            "mapping_candidates": candidates,
        })
    return {
        "template": str(path),
        "sheets": inspected,
        "recommended_manifest_reimbursement_form": {
            "template": str(path),
            "sheet": recommended_sheet,
            "output_name": "",
            "cells": recommended_mapping,
        },
    }


def main():
    args = parse_args()
    template = Path(args.template).expanduser().resolve()
    if not template.exists():
        raise SystemExit(f"Template not found: {template}")
    result = inspect_workbook(template, args.sheet, args.max_cells)
    text = json.dumps(result, ensure_ascii=False, indent=2)
    if args.out:
        output = Path(args.out).expanduser().resolve()
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(text + "\n", encoding="utf-8")
    else:
        print(text)
    if args.cells_out:
        cells_output = Path(args.cells_out).expanduser().resolve()
        cells_output.parent.mkdir(parents=True, exist_ok=True)
        form = result.get("recommended_manifest_reimbursement_form") or {}
        template_value = form.get("template", "")
        if args.relative_template and template_value:
            template_value = os.path.relpath(Path(template_value), cells_output.parent)
        reusable = {
            "template": template_value,
            "sheet": form.get("sheet", ""),
            "cells": form.get("cells", {}),
        }
        cells_output.write_text(json.dumps(reusable, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
