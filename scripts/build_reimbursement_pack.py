#!/usr/bin/env python3
import argparse
import html
import hashlib
import json
import math
import re
import shutil
import subprocess
import tempfile
import zipfile
from copy import copy
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.pagebreak import Break
from PIL import Image as PILImage
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors as rl_colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas


DEFAULT_CATEGORIES = ["交通费", "餐饮费", "设备费", "场地费", "演员费", "其他费用"]
HEIC_EXTS = {".heic", ".heif"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"} | HEIC_EXTS
PDF_EXTS = {".pdf"}
PDF_FONT = "STSong-Light"
pdfmetrics.registerFont(UnicodeCIDFont(PDF_FONT))


def parse_args():
    parser = argparse.ArgumentParser(description="Build a finance-ready reimbursement pack.")
    parser.add_argument("--manifest", required=True, help="Path to reimbursement manifest JSON.")
    parser.add_argument("--out", required=True, help="Output folder for the generated pack.")
    return parser.parse_args()


def money(value):
    return Decimal(str(value)).quantize(Decimal("0.01"))


def chinese_rmb_upper(value):
    amount = money(value)
    negative = amount < 0
    amount = abs(amount)
    fen_total = int(amount * 100)
    yuan, fraction = divmod(fen_total, 100)
    jiao, fen = divmod(fraction, 10)
    digits = "零壹贰叁肆伍陆柒捌玖"
    small_units = ("", "拾", "佰", "仟")
    group_units = ("", "万", "亿", "兆")

    def four_digits(number):
        result = ""
        pending_zero = False
        for index, divisor in enumerate((1000, 100, 10, 1)):
            digit = (number // divisor) % 10
            if digit:
                if pending_zero and result:
                    result += "零"
                result += digits[digit] + small_units[3 - index]
                pending_zero = False
            elif result and number % divisor:
                pending_zero = True
        return result

    if yuan == 0:
        integer_text = "零"
    else:
        groups = []
        remaining = yuan
        while remaining:
            groups.append(remaining % 10000)
            remaining //= 10000
        integer_text = ""
        pending_zero = False
        for index in range(len(groups) - 1, -1, -1):
            group = groups[index]
            if not group:
                if integer_text:
                    pending_zero = True
                continue
            if integer_text and (pending_zero or group < 1000):
                if not integer_text.endswith("零"):
                    integer_text += "零"
            integer_text += four_digits(group) + group_units[index]
            pending_zero = False

    result = ("负" if negative else "") + integer_text + "元"
    if jiao == 0 and fen == 0:
        return result + "整"
    if jiao:
        result += digits[jiao] + "角"
    elif fen:
        result += "零"
    if fen:
        result += digits[fen] + "分"
    return result


def safe_name(text, fallback="item"):
    text = str(text or "").strip()
    text = re.sub(r"[\\/:*?\"<>|]+", "_", text)
    text = re.sub(r"\s+", "", text)
    text = text.strip("._-")
    return text or fallback


def resolve_path(value, base_dir):
    if not value:
        return None
    path = Path(value).expanduser()
    if not path.is_absolute():
        path = base_dir / path
    return path.resolve()


def sort_key(entry):
    date_text = entry.get("date") or "9999-12-31"
    time_text = entry.get("time") or "23:59:59"
    source = entry.get("screenshot") or entry.get("source") or ""
    return (date_text, time_text, source)


def load_manifest(path):
    manifest_path = Path(path).expanduser().resolve()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    base_dir = manifest_path.parent
    top_level_invoices = []
    for invoice in manifest.get("invoices", []) or []:
        invoice_path = resolve_path(invoice, base_dir)
        if invoice_path:
            top_level_invoices.append(invoice_path)
    manifest["invoice_paths"] = top_level_invoices
    supporting_documents = []
    for document in manifest.get("supporting_documents", []) or []:
        document_path = resolve_path(document, base_dir)
        if document_path:
            supporting_documents.append(document_path)
    manifest["supporting_document_paths"] = supporting_documents
    entries = []
    for idx, raw in enumerate(manifest.get("entries", []), start=1):
        entry = dict(raw)
        entry["_source_index"] = idx
        entry["amount"] = float(money(entry["amount"]))
        screenshot = resolve_path(entry.get("screenshot"), base_dir)
        if screenshot:
            entry["screenshot_path"] = screenshot
        invoices = []
        for invoice in entry.get("invoices", []) or []:
            invoice_path = resolve_path(invoice, base_dir)
            if invoice_path:
                invoices.append(invoice_path)
        entry["invoice_paths"] = invoices
        entries.append(entry)
    entries.sort(key=sort_key)
    seen_voucher_ids = set()
    for index, entry in enumerate(entries, start=1):
        voucher_id = str(entry.get("voucher_id", "")).strip()
        if not re.fullmatch(r"FY\d{3,}", voucher_id) or voucher_id in seen_voucher_ids:
            voucher_id = f"FY{index:03d}"
        entry["voucher_id"] = voucher_id
        seen_voucher_ids.add(voucher_id)
    manifest["entries"] = entries
    form = dict(manifest.get("reimbursement_form") or {})
    form_template = resolve_path(form.get("template"), base_dir)
    if form_template:
        form["template_path"] = form_template
    manifest["reimbursement_form"] = form
    manifest["_manifest_path"] = manifest_path
    return manifest


def validate_manifest(manifest):
    errors = []
    if not manifest.get("project_name"):
        errors.append("manifest.project_name is required")
    if not manifest.get("reimburser"):
        errors.append("manifest.reimburser is required")
    voucher_ids = []
    for i, entry in enumerate(manifest.get("entries", []), start=1):
        for field in ["date", "purpose", "expense_type", "amount"]:
            if entry.get(field) in (None, ""):
                errors.append(f"entry {i}: {field} is required")
        if entry.get("screenshot_path") and not entry["screenshot_path"].exists():
            errors.append(f"entry {i}: screenshot not found: {entry['screenshot_path']}")
        for invoice in entry.get("invoice_paths", []):
            if not invoice.exists():
                errors.append(f"entry {i}: invoice not found: {invoice}")
        voucher_id = str(entry.get("voucher_id", "")).strip()
        if not re.fullmatch(r"FY\d{3,}", voucher_id):
            errors.append(f"entry {i}: invalid voucher_id: {voucher_id!r}")
        voucher_ids.append(voucher_id)
    if len(voucher_ids) != len(set(voucher_ids)):
        errors.append("voucher_id values must be unique")
    unresolved_exceptions = [
        item
        for item in manifest.get("expense_exceptions", []) or []
        if item.get("severity") == "blocking" and not item.get("resolved")
    ]
    if unresolved_exceptions:
        errors.append(f"{len(unresolved_exceptions)} blocking expense exception(s) require resolution")
    form = manifest.get("reimbursement_form") or {}
    form_template = form.get("template_path")
    if form_template and not form_template.exists():
        errors.append(f"reimbursement_form.template not found: {form_template}")
    for invoice in manifest.get("invoice_paths", []):
        if not invoice.exists():
            errors.append(f"manifest invoice not found: {invoice}")
    for document in manifest.get("supporting_document_paths", []):
        if not document.exists():
            errors.append(f"supporting document not found: {document}")
    requirement = manifest.get("invoice_requirement") or {}
    coverage = manifest.get("invoice_coverage") or {}
    policy = requirement.get("policy") or "full_amount"
    if policy == "full_amount" and coverage.get("status") != "covered":
        errors.append(
            "invoice coverage is not complete; run analyze_invoice_coverage.py and provide approved replacement invoices before final packaging"
        )
    approval_requirement = manifest.get("approval_requirement") or {}
    if approval_requirement.get("policy", "none") == "required":
        approval = manifest.get("approval_metadata") or {}
        for field in approval_requirement.get("required_fields", []) or []:
            if approval.get(field) in (None, ""):
                errors.append(f"approval_metadata.{field} is required before final packaging")
        total = sum((money(entry["amount"]) for entry in manifest.get("entries", [])), Decimal("0.00"))
        try:
            approved_amount = money(approval.get("approved_amount"))
        except Exception:
            approved_amount = None
        if approved_amount != total:
            errors.append(f"DingTalk approved amount must equal reimbursement total: {approved_amount} vs {total}")
        confirmation = manifest.get("expense_confirmation") or {}
        if confirmation.get("status") != "confirmed_by_dingtalk" or money(confirmation.get("confirmed_amount", 0)) != total:
            errors.append("expense_confirmation must confirm the exact reimbursement total through DingTalk")
    if errors:
        raise SystemExit("\n".join(errors))


def ensure_dirs(root):
    dirs = {
        "root": root,
        "tables": root / "报销表",
        "print": root / "打印",
        "finance": root / "财务提交文件夹",
        "converted": root / "_converted_images",
    }
    dirs["finance_tables"] = dirs["finance"] / "01_报销表"
    dirs["finance_screenshots"] = dirs["finance"] / "02_费用截图"
    dirs["finance_invoices"] = dirs["finance"] / "03_发票"
    dirs["finance_print"] = dirs["finance"] / "04_打印文件"
    for folder in dirs.values():
        folder.mkdir(parents=True, exist_ok=True)
    return dirs


def converted_image_name(image_path):
    digest = hashlib.sha1(str(image_path).encode("utf-8")).hexdigest()[:10]
    return f"{safe_name(image_path.stem)}_{digest}.png"


def convert_heic_image(image_path, dirs):
    output = dirs["converted"] / converted_image_name(image_path)
    if output.exists():
        return output
    pillow_error = ""
    try:
        from pillow_heif import register_heif_opener

        register_heif_opener()
        with PILImage.open(image_path) as image:
            image.convert("RGB").save(output, "PNG")
        if output.exists():
            return output
    except Exception as exc:
        pillow_error = str(exc)

    sips = shutil.which("sips")
    if not sips:
        raise SystemExit(
            f"HEIC/HEIF conversion failed for {image_path}. Install pillow-heif on Windows/macOS. "
            f"macOS can also use sips. Pillow error: {pillow_error or 'pillow-heif unavailable'}"
        )
    result = subprocess.run(
        [sips, "-s", "format", "png", str(image_path), "--out", str(output)],
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
    )
    if result.returncode != 0 or not output.exists():
        raise SystemExit(
            f"Failed to convert HEIC/HEIF image to PNG: {image_path}\n"
            f"Pillow: {pillow_error or 'not used'}\nsips: {result.stderr.strip()}"
        )
    return output


def prepare_image_for_output(image_path, dirs):
    if image_path and image_path.suffix.lower() in HEIC_EXTS:
        return convert_heic_image(image_path, dirs)
    return image_path


def prepare_manifest_images(manifest, dirs):
    for entry in manifest["entries"]:
        if entry.get("screenshot_path"):
            entry["screenshot_path"] = prepare_image_for_output(entry["screenshot_path"], dirs)
        entry["invoice_paths"] = [prepare_image_for_output(invoice, dirs) for invoice in entry.get("invoice_paths", [])]
    manifest["invoice_paths"] = [prepare_image_for_output(invoice, dirs) for invoice in manifest.get("invoice_paths", [])]
    manifest["supporting_document_paths"] = [
        prepare_image_for_output(document, dirs)
        for document in manifest.get("supporting_document_paths", [])
    ]


def categories_for(manifest):
    categories = list(manifest.get("categories") or DEFAULT_CATEGORIES)
    for entry in manifest["entries"]:
        if entry["expense_type"] not in categories:
            categories.append(entry["expense_type"])
    return categories


def make_date(date_text):
    return datetime.strptime(date_text, "%Y-%m-%d").date()


def copy_cell_style(source, target):
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format


def fit_image_size(image_path, max_width, max_height):
    with PILImage.open(image_path) as img:
        width, height = img.size
    scale = min(max_width / width, max_height / height)
    return max(1, int(width * scale)), max(1, int(height * scale))


def add_xl_image(ws, image_path, cell, max_width, max_height):
    img = XLImage(str(image_path))
    img.width, img.height = fit_image_size(image_path, max_width, max_height)
    ws.add_image(img, cell)


def build_reimbursement_workbook(manifest, dirs):
    entries = manifest["entries"]
    categories = categories_for(manifest)
    project = manifest["project_name"]
    reimburser = manifest["reimburser"]
    summary_start = 3
    summary_end = summary_start + len(categories) - 1
    summary_total_row = summary_end + 1
    header_row = summary_total_row + 1
    detail_start = header_row + 1
    detail_end = detail_start + len(entries) - 1
    total_row = detail_end + 1

    wb = Workbook()
    ws = wb.active
    ws.title = "报销明细"
    ws.sheet_view.showGridLines = False

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)
    for row in range(summary_start, summary_total_row + 1):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
        ws.merge_cells(start_row=row, start_column=7, end_row=row, end_column=8)
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=5)
    ws.merge_cells(start_row=total_row, start_column=7, end_row=total_row, end_column=8)

    ws["A1"] = f"{project}{reimburser} 报销明细表"
    ws["A2"] = "分类统计"
    category_totals = {category: Decimal("0.00") for category in categories}
    category_counts = {category: 0 for category in categories}
    for entry in entries:
        category = entry["expense_type"]
        category_totals[category] = category_totals.get(category, Decimal("0.00")) + money(entry["amount"])
        category_counts[category] = category_counts.get(category, 0) + 1
    for offset, category in enumerate(categories):
        row = summary_start + offset
        ws[f"A{row}"] = category
        ws[f"E{row}"] = float(category_totals.get(category, Decimal("0.00")))
        ws[f"G{row}"] = f"({category_counts.get(category, 0)}笔)"
    ws[f"A{summary_total_row}"] = "总  计"
    ws[f"E{summary_total_row}"] = float(sum(category_totals.values(), Decimal("0.00")))
    ws[f"G{summary_total_row}"] = f"({len(entries)}笔)"
    ws.append([])
    for col, value in enumerate(["凭证编号", "日期", "用途", "报销人", "费用类型", "金额(元)", "凭证截图", "发票"], start=1):
        ws.cell(header_row, col, value)

    for row_index, entry in enumerate(entries, start=detail_start):
        ws.cell(row_index, 1, entry["voucher_id"])
        ws.cell(row_index, 2, make_date(entry["date"]))
        ws.cell(row_index, 3, entry["purpose"])
        ws.cell(row_index, 4, reimburser)
        ws.cell(row_index, 5, entry["expense_type"])
        ws.cell(row_index, 6, entry["amount"])
        if entry.get("invoice_paths"):
            ws.cell(row_index, 8, "；".join(p.name for p in entry["invoice_paths"]))
        screenshot = entry.get("screenshot_path")
        if screenshot:
            add_xl_image(ws, screenshot, f"G{row_index}", 250, 520)

    ws.cell(total_row, 1, "合计")
    ws.cell(total_row, 6, float(sum((money(e["amount"]) for e in entries), Decimal("0.00"))))

    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    yellow = PatternFill("solid", fgColor="FFFF00")
    blue = PatternFill("solid", fgColor="D9E1F2")

    for row in ws.iter_rows(min_row=1, max_row=total_row, min_col=1, max_col=8):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
            cell.border = border
    ws["A1"].font = Font(name="Arial", size=12, bold=True)
    for row in range(2, header_row + 1):
        for cell in ws[row]:
            if cell.column <= 8:
                cell.font = Font(name="Arial", size=10, bold=True)
                cell.fill = blue if row == 2 else yellow
    for cell in ws[total_row]:
        if cell.column <= 8:
            cell.font = Font(name="Arial", size=10, bold=True)
            cell.fill = yellow
    for row in range(detail_start, total_row):
        ws.cell(row, 3).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = {"A": 6, "B": 14, "C": 32, "D": 13.1, "E": 10, "F": 12, "G": 60, "H": 12}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in range(1, summary_total_row + 1):
        ws.row_dimensions[row].height = 35
    ws.row_dimensions[header_row].height = 25
    for row in range(detail_start, total_row):
        ws.row_dimensions[row].height = 409.5
    ws.row_dimensions[total_row].height = 28
    for row in range(summary_start, summary_total_row + 1):
        ws[f"E{row}"].number_format = "#,##0.00"
    for row in range(detail_start, total_row + 1):
        ws[f"B{row}"].number_format = "yyyy-mm-dd"
        ws[f"F{row}"].number_format = "#,##0.00"

    output = dirs["tables"] / f"{safe_name(project)}_{safe_name(reimburser)}_报销明细表.xlsx"
    wb.save(output)
    return output


def manifest_context(manifest):
    category_totals = {category: Decimal("0.00") for category in categories_for(manifest)}
    category_counts = {category: 0 for category in categories_for(manifest)}
    total = Decimal("0.00")
    dates = []
    for entry in manifest["entries"]:
        amount = money(entry["amount"])
        total += amount
        category = entry["expense_type"]
        category_totals[category] = category_totals.get(category, Decimal("0.00")) + amount
        category_counts[category] = category_counts.get(category, 0) + 1
        dates.append(entry["date"])
    approval = manifest.get("approval_metadata") or {}
    context = {
        "project_name": manifest["project_name"],
        "reimburser": manifest["reimburser"],
        "dingding_number": approval.get("dingding_number", ""),
        "reimbursement_type": approval.get("reimbursement_type", ""),
        "payee": approval.get("payee", ""),
        "invoice_entity": approval.get("invoice_entity", ""),
        "approval_company": approval.get("company", ""),
        "reimbursement_project": approval.get("reimbursement_project", ""),
        "contract_number": approval.get("contract_number", ""),
        "reimbursement_reason": approval.get("reimbursement_reason", ""),
        "business_line": approval.get("business_line", ""),
        "approved_amount": f"{money(approval.get('approved_amount')):.2f}" if approval.get("approved_amount") not in (None, "") else "",
        "expense_count": len(manifest["entries"]),
        "total_amount": f"{total:.2f}",
        "date_start": min(dates) if dates else "",
        "date_end": max(dates) if dates else "",
        "date_range": f"{min(dates)}至{max(dates)}" if dates else "",
    }
    screenshot_count = sum(1 for entry in manifest["entries"] if entry.get("screenshot_path"))
    screenshot_print_pages = math.ceil(screenshot_count / 3) if screenshot_count else 0
    invoice_print_pages = sum(print_document_page_count(path) for path in collect_invoice_print_documents(manifest))
    form = manifest.get("reimbursement_form") or {}
    expense_form_print_pages = manifest.get("_expense_form_print_pages")
    if expense_form_print_pages is None:
        expense_form_print_pages = int(form.get("print_pages", 1)) if form.get("template_path") else 0
    context["screenshot_print_pages"] = screenshot_print_pages
    context["invoice_print_pages"] = invoice_print_pages
    context["expense_form_print_pages"] = expense_form_print_pages
    context["total_print_pages"] = expense_form_print_pages + screenshot_print_pages + invoice_print_pages
    for category in categories_for(manifest):
        amount = category_totals.get(category, Decimal("0.00"))
        key = safe_name(category)
        context[f"{key}_amount"] = f"{amount:.2f}"
        context[f"{key}_count"] = category_counts.get(category, 0)
    return context


def render_cell_value(value, context):
    if isinstance(value, str):
        def replace(match):
            key = match.group(1).strip()
            return str(context.get(key, match.group(0)))
        rendered = re.sub(r"\{\{\s*([^{}]+?)\s*\}\}", replace, value)
        if re.fullmatch(r"-?\d+(?:\.\d+)?", rendered):
            integer_part = rendered.lstrip("-").split(".", 1)[0]
            if len(integer_part) > 15:
                return rendered
            return float(rendered) if "." in rendered else int(rendered)
        return rendered
    return value


def writable_cell_ref(ws, cell_ref):
    for cell_range in ws.merged_cells.ranges:
        if cell_ref in cell_range:
            return f"{get_column_letter(cell_range.min_col)}{cell_range.min_row}"
    return cell_ref


def formula_cache_values(ws, context):
    caches = {}
    sum_pattern = re.compile(r"^=?SUM\(\s*([A-Z]{1,3}\d+):([A-Z]{1,3}\d+)\s*\)$", re.IGNORECASE)
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type != "f" or not isinstance(cell.value, str):
                continue
            formula = cell.value.strip()
            match = sum_pattern.fullmatch(formula)
            if match:
                total = Decimal("0.00")
                for range_row in ws[f"{match.group(1)}:{match.group(2)}"]:
                    for source_cell in range_row:
                        if isinstance(source_cell.value, (int, float, Decimal)) and not isinstance(source_cell.value, bool):
                            total += money(source_cell.value)
                caches[cell.coordinate] = (f"{money(total):.2f}", None)
            elif "[DBNUM2]" in formula.upper():
                caches[cell.coordinate] = (chinese_rmb_upper(context["total_amount"]), "str")
    return caches


def patch_formula_caches(path, sheet_index, caches):
    if not caches:
        return
    worksheet_name = f"xl/worksheets/sheet{sheet_index}.xml"
    with zipfile.ZipFile(path, "r") as source:
        if worksheet_name not in source.namelist():
            raise ValueError(f"Worksheet XML not found for formula cache refresh: {worksheet_name}")
        handle = tempfile.NamedTemporaryFile(prefix="expense-form-", suffix=path.suffix, dir=path.parent, delete=False)
        temp_path = Path(handle.name)
        handle.close()
        try:
            with zipfile.ZipFile(temp_path, "w") as target:
                for item in source.infolist():
                    data = source.read(item.filename)
                    if item.filename == worksheet_name:
                        xml = data.decode("utf-8")
                        for cell_ref, (cached_value, cached_type) in caches.items():
                            pattern = re.compile(
                                rf'(<c r="{re.escape(cell_ref)}"[^>]*>)(.*?)(</c>)',
                                re.DOTALL,
                            )

                            def replace_cell(match):
                                header, body, footer = match.groups()
                                header = re.sub(r'\s+t="[^"]*"', "", header)
                                if cached_type:
                                    header = header[:-1] + f' t="{cached_type}">'
                                escaped = html.escape(str(cached_value))
                                if re.search(r"<v>.*?</v>", body, flags=re.DOTALL):
                                    body = re.sub(r"<v>.*?</v>", f"<v>{escaped}</v>", body, count=1, flags=re.DOTALL)
                                else:
                                    body += f"<v>{escaped}</v>"
                                return header + body + footer

                            xml, replaced = pattern.subn(replace_cell, xml, count=1)
                            if replaced != 1:
                                raise ValueError(f"Formula cell not found while refreshing cache: {cell_ref}")
                        data = xml.encode("utf-8")
                    target.writestr(item, data)
            temp_path.replace(path)
        except Exception:
            temp_path.unlink(missing_ok=True)
            raise


def build_expense_form(manifest, dirs):
    form = manifest.get("reimbursement_form") or {}
    template = form.get("template_path")
    if not template:
        return None
    output_name = form.get("output_name") or f"{safe_name(manifest['project_name'])}_{safe_name(manifest['reimburser'])}_费用报销单{template.suffix}"
    output = dirs["tables"] / safe_name(output_name, "费用报销单.xlsx")
    if template.suffix.lower() not in {".xlsx", ".xlsm"}:
        shutil.copy2(template, output)
        return output

    wb = load_workbook(template, keep_vba=template.suffix.lower() == ".xlsm")
    sheet_name = form.get("sheet")
    ws = wb[sheet_name] if sheet_name else wb.active
    sheet_index = wb.worksheets.index(ws) + 1
    context = manifest_context(manifest)
    for cell_ref, value in (form.get("cells") or {}).items():
        ws[writable_cell_ref(ws, cell_ref)] = render_cell_value(value, context)
    category_rows = form.get("category_rows") or []
    if category_rows:
        nonzero_categories = []
        for category in categories_for(manifest):
            amount = money(context.get(f"{safe_name(category)}_amount", "0.00"))
            if amount != Decimal("0.00"):
                nonzero_categories.append((category, amount))
        if len(nonzero_categories) > len(category_rows):
            raise SystemExit(
                f"费用报销单只有 {len(category_rows)} 行费用分类空间，但本项目有 {len(nonzero_categories)} 个非零分类；请扩充模板后再生成。"
            )
        for index, row_config in enumerate(category_rows):
            label_cell = row_config.get("label_cell")
            amount_cell = row_config.get("amount_cell")
            if not label_cell or not amount_cell:
                raise SystemExit("reimbursement_form.category_rows 每项都必须包含 label_cell 和 amount_cell。")
            label_target = writable_cell_ref(ws, label_cell)
            amount_target = writable_cell_ref(ws, amount_cell)
            if index < len(nonzero_categories):
                category, amount = nonzero_categories[index]
                ws[label_target] = category
                ws[amount_target] = float(amount)
            else:
                ws[label_target] = None
                ws[amount_target] = None
    caches = formula_cache_values(ws, context)
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    output = output.with_suffix(template.suffix)
    wb.save(output)
    patch_formula_caches(output, sheet_index, caches)
    return output


def form_print_area(ws):
    if ws.print_area:
        area = str(ws.print_area).split("!", 1)[-1].replace("$", "").strip("'")
        return range_boundaries(area)
    return 1, 1, ws.max_column, ws.max_row


def excel_column_points(ws, column):
    width = ws.column_dimensions[get_column_letter(column)].width
    if width is None:
        width = ws.sheet_format.defaultColWidth or 8.43
    return max(4.0, float(width) * 7.0)


def excel_row_points(ws, row):
    height = ws.row_dimensions[row].height
    if height is None:
        height = ws.sheet_format.defaultRowHeight or 15.0
    return max(4.0, float(height))


def format_form_value(value, number_format):
    if value is None:
        return ""
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        number = float(value)
        if "0.00" in (number_format or ""):
            text = f"{number:.2f}"
        else:
            text = f"{number:g}"
        if "¥" in (number_format or "") or "￥" in (number_format or ""):
            text = f"¥{text}"
        return text
    return str(value)


def wrapped_form_lines(text, font_size, max_width, wrap):
    source_lines = text.splitlines() or [""]
    if not wrap:
        return source_lines
    lines = []
    for source in source_lines:
        current = ""
        for char in source:
            candidate = current + char
            if current and pdfmetrics.stringWidth(candidate, PDF_FONT, font_size) > max_width:
                lines.append(current)
                current = char
            else:
                current = candidate
        lines.append(current)
    return lines


def form_border_width(style):
    if style in {"medium", "mediumDashed", "mediumDashDot", "mediumDashDotDot"}:
        return 1.4
    if style in {"thick"}:
        return 2.0
    return 0.75


def draw_form_border(c, side, x1, y1, x2, y2, scale):
    if not side or not side.style:
        return
    c.setStrokeColor(rl_colors.black)
    c.setLineWidth(form_border_width(side.style) * min(scale, 1.5))
    c.line(x1, y1, x2, y2)


def render_xlsx_form_pdf(expense_form, output, sheet_name=None):
    style_wb = load_workbook(expense_form, keep_vba=expense_form.suffix.lower() == ".xlsm", data_only=False)
    value_wb = load_workbook(expense_form, keep_vba=expense_form.suffix.lower() == ".xlsm", data_only=True)
    style_ws = style_wb[sheet_name] if sheet_name else style_wb.active
    value_ws = value_wb[style_ws.title]
    min_col, min_row, max_col, max_row = form_print_area(style_ws)

    page_size = landscape(A4) if style_ws.page_setup.orientation != "portrait" else A4
    page_w, page_h = page_size
    margin = 22.0
    column_widths = {col: excel_column_points(style_ws, col) for col in range(min_col, max_col + 1)}
    row_heights = {row: excel_row_points(style_ws, row) for row in range(min_row, max_row + 1)}
    raw_width = sum(column_widths.values())
    raw_height = sum(row_heights.values())
    scale = min((page_w - 2 * margin) / raw_width, (page_h - 2 * margin) / raw_height)
    content_w = raw_width * scale
    content_h = raw_height * scale
    left = (page_w - content_w) / 2
    top = (page_h + content_h) / 2

    x_positions = {min_col: left}
    for col in range(min_col, max_col + 1):
        x_positions[col + 1] = x_positions[col] + column_widths[col] * scale
    y_positions = {min_row: top}
    for row in range(min_row, max_row + 1):
        y_positions[row + 1] = y_positions[row] - row_heights[row] * scale

    merged_top_left = {}
    merged_covered = set()
    for merged in style_ws.merged_cells.ranges:
        if merged.max_col < min_col or merged.min_col > max_col or merged.max_row < min_row or merged.min_row > max_row:
            continue
        key = (merged.min_row, merged.min_col)
        merged_top_left[key] = (merged.min_row, merged.min_col, merged.max_row, merged.max_col)
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                if (row, col) != key:
                    merged_covered.add((row, col))

    c = canvas.Canvas(str(output), pagesize=page_size)
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            if (row, col) in merged_covered:
                continue
            min_r, min_c, max_r, max_c = merged_top_left.get((row, col), (row, col, row, col))
            x1 = x_positions[min_c]
            x2 = x_positions[max_c + 1]
            y1 = y_positions[max_r + 1]
            y2 = y_positions[min_r]
            cell = style_ws.cell(min_r, min_c)

            if cell.fill and cell.fill.fill_type == "solid":
                color = cell.fill.fgColor
                if color.type == "rgb" and color.rgb:
                    rgb = color.rgb[-6:]
                    c.setFillColorRGB(*(int(rgb[index:index + 2], 16) / 255 for index in (0, 2, 4)))
                else:
                    c.setFillColor(rl_colors.HexColor("#F2F2F2"))
                c.rect(x1, y1, x2 - x1, y2 - y1, stroke=0, fill=1)

            draw_form_border(c, cell.border.left, x1, y1, x1, y2, scale)
            draw_form_border(c, cell.border.right, x2, y1, x2, y2, scale)
            draw_form_border(c, cell.border.top, x1, y2, x2, y2, scale)
            draw_form_border(c, cell.border.bottom, x1, y1, x2, y1, scale)

            raw_value = value_ws.cell(min_r, min_c).value if cell.data_type == "f" else cell.value
            text = format_form_value(raw_value, cell.number_format)
            if not text.strip():
                continue
            padding = 3.0 * scale
            max_text_width = max(1.0, x2 - x1 - 2 * padding)
            font_size = max(6.0, float(cell.font.sz or 10.0) * scale)
            digit_identifier = bool(re.fullmatch(r"\d{16,}", text))
            if digit_identifier:
                while font_size > 6.0 and pdfmetrics.stringWidth(text, PDF_FONT, font_size) > max_text_width:
                    font_size -= 0.25
            lines = wrapped_form_lines(text, font_size, max_text_width, bool(cell.alignment.wrap_text and not digit_identifier))
            line_height = font_size * 1.18
            available_height = y2 - y1 - 2 * padding
            while len(lines) * line_height > available_height and font_size > 6.0:
                font_size -= 0.25
                line_height = font_size * 1.18
                lines = wrapped_form_lines(text, font_size, max_text_width, bool(cell.alignment.wrap_text and not digit_identifier))

            total_text_height = len(lines) * line_height
            vertical = cell.alignment.vertical or "center"
            if vertical == "top":
                baseline = y2 - padding - font_size
            elif vertical == "bottom":
                baseline = y1 + padding + total_text_height - line_height
            else:
                baseline = y1 + (y2 - y1 + total_text_height) / 2 - line_height
            horizontal = cell.alignment.horizontal
            if not horizontal:
                horizontal = "right" if isinstance(raw_value, (int, float, Decimal)) else "left"

            c.setFillColor(rl_colors.black)
            c.setFont(PDF_FONT, font_size)
            for index, line in enumerate(lines):
                line_width = pdfmetrics.stringWidth(line, PDF_FONT, font_size)
                if horizontal in {"center", "centerContinuous"}:
                    text_x = x1 + (x2 - x1 - line_width) / 2
                elif horizontal == "right":
                    text_x = x2 - padding - line_width
                else:
                    text_x = x1 + padding
                text_y = baseline - index * line_height
                c.drawString(text_x, text_y, line)
                if cell.font.bold:
                    c.drawString(text_x + 0.18, text_y, line)
                if cell.font.underline:
                    c.setLineWidth(0.6)
                    c.line(text_x, text_y - 1.2, text_x + line_width, text_y - 1.2)
    c.showPage()
    c.save()


def build_expense_form_pdf(manifest, dirs, expense_form):
    if not expense_form:
        return None
    if expense_form.suffix.lower() == ".pdf":
        return expense_form
    if expense_form.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise SystemExit(f"费用报销单无法转换为打印 PDF: {expense_form}")
    output = dirs["print"] / f"{safe_name(manifest['project_name'])}_费用报销单_打印.pdf"
    form = manifest.get("reimbursement_form") or {}
    render_xlsx_form_pdf(expense_form, output, form.get("sheet"))
    return output


def draw_fit_image(c, image_path, box_x, box_y, box_w, box_h):
    with PILImage.open(image_path) as img:
        width, height = img.size
    scale = min(box_w / width, box_h / height)
    draw_w = width * scale
    draw_h = height * scale
    x = box_x + (box_w - draw_w) / 2
    y = box_y + (box_h - draw_h) / 2
    c.drawImage(ImageReader(str(image_path)), x, y, width=draw_w, height=draw_h, preserveAspectRatio=True, mask="auto")


def build_screenshot_pdf(manifest, dirs):
    entries = [e for e in manifest["entries"] if e.get("screenshot_path")]
    if not entries:
        return None
    project = manifest["project_name"]
    output = dirs["print"] / f"{safe_name(project)}_费用截图_三张一页.pdf"
    page_w, page_h = landscape(A4)
    margin = 28
    gap = 12
    caption_h = 30
    slot_w = (page_w - 2 * margin - 2 * gap) / 3
    image_h = page_h - 2 * margin - caption_h

    c = canvas.Canvas(str(output), pagesize=landscape(A4))
    for page_index in range(0, len(entries), 3):
        page_entries = entries[page_index:page_index + 3]
        c.setFont(PDF_FONT, 10)
        c.drawString(margin, page_h - margin + 4, f"{project} 费用截图 第 {page_index // 3 + 1} 页")
        for col, entry in enumerate(page_entries):
            x = margin + col * (slot_w + gap)
            caption_y = page_h - margin - caption_h + 8
            caption = f"{entry['voucher_id']} {entry['date']} {entry['purpose']} ¥{money(entry['amount'])}"
            c.setFont(PDF_FONT, 8)
            c.drawString(x, caption_y, caption[:46])
            draw_fit_image(c, entry["screenshot_path"], x, margin, slot_w, image_h)
            c.rect(x, margin, slot_w, image_h + caption_h, stroke=1, fill=0)
        c.showPage()
    c.save()
    return output


def pdf_from_image(image_path, output_pdf):
    c = canvas.Canvas(str(output_pdf), pagesize=A4)
    page_w, page_h = A4
    margin = 24
    draw_fit_image(c, image_path, margin, margin, page_w - 2 * margin, page_h - 2 * margin)
    c.showPage()
    c.save()


def collect_invoices(manifest):
    invoices = []
    seen = set()
    for invoice in manifest.get("invoice_paths", []):
        key = str(invoice)
        if key not in seen:
            seen.add(key)
            invoices.append(invoice)
    for entry in manifest["entries"]:
        for invoice in entry.get("invoice_paths", []):
            key = str(invoice)
            if key not in seen:
                seen.add(key)
                invoices.append(invoice)
    return invoices


def collect_invoice_print_documents(manifest):
    documents = collect_invoices(manifest)
    seen = {str(path) for path in documents}
    for document in manifest.get("supporting_document_paths", []):
        key = str(document)
        if key in seen:
            continue
        seen.add(key)
        documents.append(document)
    return documents


def print_document_page_count(path):
    suffix = path.suffix.lower()
    if suffix in PDF_EXTS:
        return len(PdfReader(str(path)).pages)
    if suffix in IMAGE_EXTS:
        return 1
    return 0


def build_invoice_pdf(manifest, dirs):
    documents = collect_invoice_print_documents(manifest)
    if not documents:
        return None
    output = dirs["print"] / f"{safe_name(manifest['project_name'])}_发票_一张一页.pdf"
    writer = PdfWriter()
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_dir = Path(tmp_dir)
        for index, document in enumerate(documents, start=1):
            suffix = document.suffix.lower()
            if suffix in PDF_EXTS:
                reader = PdfReader(str(document))
                for page in reader.pages:
                    writer.add_page(page)
            elif suffix in IMAGE_EXTS:
                temp_pdf = tmp_dir / f"invoice_{index}.pdf"
                pdf_from_image(document, temp_pdf)
                reader = PdfReader(str(temp_pdf))
                writer.add_page(reader.pages[0])
        with output.open("wb") as fh:
            writer.write(fh)
    return output


def append_pdf(writer, pdf_path):
    if not pdf_path:
        return 0
    reader = PdfReader(str(pdf_path))
    for page in reader.pages:
        writer.add_page(page)
    return len(reader.pages)


def pdf_page_count(pdf_path):
    return len(PdfReader(str(pdf_path)).pages)


def build_combined_print_pdf(manifest, dirs, expense_form_pdf, screenshot_pdf, invoice_pdf):
    if not expense_form_pdf and not screenshot_pdf and not invoice_pdf:
        return None
    output = dirs["print"] / f"{safe_name(manifest['project_name'])}_全部打印_一键打印.pdf"
    writer = PdfWriter()
    append_pdf(writer, expense_form_pdf)
    append_pdf(writer, screenshot_pdf)
    append_pdf(writer, invoice_pdf)
    total_pages = len(writer.pages)
    for page_index, page in enumerate(writer.pages, start=1):
        page_w = float(page.mediabox.width)
        page_h = float(page.mediabox.height)
        overlay_buffer = BytesIO()
        overlay_canvas = canvas.Canvas(overlay_buffer, pagesize=(page_w, page_h))
        overlay_canvas.setFillColor(rl_colors.HexColor("#555555"))
        overlay_canvas.setFont("Helvetica", 7)
        overlay_canvas.drawRightString(page_w - 14, 10, f"{page_index} / {total_pages}")
        overlay_canvas.showPage()
        overlay_canvas.save()
        overlay_buffer.seek(0)
        overlay_page = PdfReader(overlay_buffer).pages[0]
        page.merge_page(overlay_page)
    with output.open("wb") as fh:
        writer.write(fh)
    return output


def build_print_workbook(manifest, dirs):
    entries = [e for e in manifest["entries"] if e.get("screenshot_path")]
    invoices = collect_invoice_print_documents(manifest)
    project = manifest["project_name"]
    wb = Workbook()
    ws = wb.active
    ws.title = "费用截图打印"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    for col in ["A", "B", "C"]:
        ws.column_dimensions[col].width = 35

    row = 1
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for page_index in range(0, len(entries), 3):
        page_entries = entries[page_index:page_index + 3]
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        ws.cell(row, 1, f"{project} 费用截图 第 {page_index // 3 + 1} 页")
        ws.cell(row, 1).font = Font(name="Arial", size=12, bold=True)
        ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 24
        caption_row = row + 1
        image_row = row + 2
        ws.row_dimensions[caption_row].height = 24
        ws.row_dimensions[image_row].height = 370
        for col_index, entry in enumerate(page_entries, start=1):
            col = get_column_letter(col_index)
            caption = f"{entry['voucher_id']} {entry['date']} {entry['purpose']} ¥{money(entry['amount'])}"
            ws.cell(caption_row, col_index, caption)
            ws.cell(caption_row, col_index).alignment = Alignment(horizontal="center", vertical="center")
            for target_row in [caption_row, image_row]:
                ws.cell(target_row, col_index).border = border
            add_xl_image(ws, entry["screenshot_path"], f"{col}{image_row}", 230, 490)
        ws.row_breaks.append(Break(id=image_row + 1))
        row += 4

    invoice_ws = wb.create_sheet("发票打印索引")
    invoice_ws.append(["序号", "发票文件", "打印说明"])
    for i, invoice in enumerate(invoices, start=1):
        invoice_ws.append([i, invoice.name, "PDF 发票请使用生成的发票打印 PDF；图片发票也已合入 PDF。"])
    for col, width in {"A": 8, "B": 50, "C": 70}.items():
        invoice_ws.column_dimensions[col].width = width
    for cell in invoice_ws[1]:
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
    output = dirs["print"] / f"{safe_name(project)}_打印排版.xlsx"
    wb.save(output)
    return output


def invoice_items_by_path(manifest):
    base_dir = manifest.get("_manifest_path", Path.cwd()).parent
    items = {}
    for raw in manifest.get("invoice_items", []) or []:
        if not isinstance(raw, dict) or not raw.get("path"):
            continue
        path = resolve_path(raw["path"], base_dir)
        if path:
            items[str(path)] = raw
    return items


def copy_sources_and_outputs(manifest, dirs, generated):
    project = safe_name(manifest["project_name"])
    copied_screenshots = []
    copied_invoices = []
    copied_supporting_documents = []
    invoice_items = invoice_items_by_path(manifest)
    for index, entry in enumerate(manifest["entries"], start=1):
        prefix = f"{entry['voucher_id']}_{entry['date']}_{safe_name(entry['purpose'])}_{money(entry['amount'])}"
        screenshot = entry.get("screenshot_path")
        if screenshot:
            dest = dirs["finance_screenshots"] / f"{prefix}_{safe_name(screenshot.stem)}{screenshot.suffix.lower()}"
            shutil.copy2(screenshot, dest)
            copied_screenshots.append(dest)
        for invoice_index, invoice in enumerate(entry.get("invoice_paths", []), start=1):
            dest = dirs["finance_invoices"] / f"{prefix}_发票{invoice_index}_{safe_name(invoice.stem)}{invoice.suffix.lower()}"
            shutil.copy2(invoice, dest)
            copied_invoices.append(dest)
    for invoice_index, invoice in enumerate(manifest.get("invoice_paths", []), start=1):
        item = invoice_items.get(str(invoice), {})
        role_label = "替代发票" if item.get("role") == "replacement" else "原始发票"
        amount_label = f"_{money(item['amount'])}" if item.get("amount") not in (None, "") else ""
        voucher_label = "_".join(item.get("voucher_ids", []) or []) or f"{invoice_index:03d}"
        dest = dirs["finance_invoices"] / f"{voucher_label}_{role_label}{amount_label}_{safe_name(invoice.stem)}{invoice.suffix.lower()}"
        shutil.copy2(invoice, dest)
        copied_invoices.append(dest)
    for document_index, document in enumerate(manifest.get("supporting_document_paths", []), start=1):
        dest = dirs["finance_invoices"] / f"{document_index:03d}_发票附件_行程单_{safe_name(document.stem)}{document.suffix.lower()}"
        shutil.copy2(document, dest)
        copied_supporting_documents.append(dest)

    for path in generated:
        if not path:
            continue
        if path.suffix.lower() in {".xlsx", ".xlsm"} and ("报销明细表" in path.name or "费用报销单" in path.name):
            shutil.copy2(path, dirs["finance_tables"] / path.name)
        else:
            shutil.copy2(path, dirs["finance_print"] / path.name)

    manifest_copy = dirs["finance"] / "manifest.json"
    cleaned = json.loads(json.dumps(manifest, default=str, ensure_ascii=False))
    cleaned.pop("_manifest_path", None)
    for entry in cleaned.get("entries", []):
        entry.pop("screenshot_path", None)
        entry.pop("invoice_paths", None)
        entry.pop("_source_index", None)
    cleaned.pop("invoice_paths", None)
    cleaned.pop("supporting_document_paths", None)
    if cleaned.get("reimbursement_form"):
        cleaned["reimbursement_form"].pop("template_path", None)
    manifest_copy.write_text(json.dumps(cleaned, ensure_ascii=False, indent=2), encoding="utf-8")

    coverage = manifest.get("invoice_coverage") or {}
    report_names = {
        "report": "invoice_coverage.json",
        "markdown_report": "发票缺口清单.md",
    }
    for key, output_name in report_names.items():
        value = coverage.get(key)
        if not value:
            continue
        source = Path(value)
        if source.exists():
            shutil.copy2(source, dirs["finance"] / output_name)
    for value in (manifest.get("workflow_reports") or {}).values():
        if not value:
            continue
        source = Path(value)
        if source.exists():
            shutil.copy2(source, dirs["finance"] / source.name)
    return copied_screenshots, copied_invoices, copied_supporting_documents


def should_skip_zip_path(path):
    parts = set(path.parts)
    return (
        path.name == ".DS_Store"
        or path.name.startswith("._")
        or "__MACOSX" in parts
        or "__pycache__" in parts
        or path.suffix == ".pyc"
    )


def build_finance_zip(manifest, dirs):
    output = dirs["root"] / f"{safe_name(manifest['project_name'])}_{safe_name(manifest['reimburser'])}_财务提交文件夹.zip"
    finance = dirs["finance"]
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for folder in sorted([p for p in finance.rglob("*") if p.is_dir()]):
            if should_skip_zip_path(folder):
                continue
            arcname = Path(finance.name) / folder.relative_to(finance)
            zf.writestr(str(arcname).rstrip("/") + "/", "")
        for path in sorted([p for p in finance.rglob("*") if p.is_file()]):
            if should_skip_zip_path(path):
                continue
            arcname = Path(finance.name) / path.relative_to(finance)
            zf.write(path, arcname)
    return output


def expected_finance_zip_path(manifest, dirs):
    return dirs["root"] / f"{safe_name(manifest['project_name'])}_{safe_name(manifest['reimburser'])}_财务提交文件夹.zip"


def path_line(path):
    return str(path) if path else "无"


def write_delivery_checklist(
    manifest,
    dirs,
    reimbursement,
    expense_form,
    expense_form_pdf,
    screenshot_pdf,
    invoice_pdf,
    combined_print_pdf,
    print_workbook,
    copied_screenshots,
    copied_invoices,
    copied_supporting_documents,
    finance_zip,
):
    total = sum((money(e["amount"]) for e in manifest["entries"]), Decimal("0.00"))
    expense_form_pages = pdf_page_count(expense_form_pdf) if expense_form_pdf and expense_form_pdf.exists() else 0
    screenshot_pages = pdf_page_count(screenshot_pdf) if screenshot_pdf and screenshot_pdf.exists() else 0
    invoice_pages = pdf_page_count(invoice_pdf) if invoice_pdf and invoice_pdf.exists() else 0
    combined_pages = pdf_page_count(combined_print_pdf) if combined_print_pdf and combined_print_pdf.exists() else 0
    categories = {category: Decimal("0.00") for category in categories_for(manifest)}
    for entry in manifest["entries"]:
        categories.setdefault(entry["expense_type"], Decimal("0.00"))
        categories[entry["expense_type"]] += money(entry["amount"])
    category_text = "\n".join(f"- {category}: {amount:.2f}" for category, amount in categories.items())
    form_status = "已生成" if expense_form else "未生成，需提供或接入真实费用报销单模板"
    invoice_status = (
        f"{len(copied_invoices)} 个发票文件，{len(copied_supporting_documents)} 个行程单附件，打印 PDF {invoice_pages} 页"
        if copied_invoices or copied_supporting_documents
        else "无发票或行程单文件"
    )
    coverage = manifest.get("invoice_coverage") or {}
    coverage_text = (
        f"应覆盖 {coverage.get('required_amount', '0.00')} 元，"
        f"已识别 {coverage.get('recognized_amount', '0.00')} 元，"
        f"缺口 {coverage.get('missing_amount', '0.00')} 元；"
        f"替代发票 {coverage.get('replacement_invoice_count', 0)} 张，"
        f"合计 {coverage.get('replacement_invoice_amount', '0.00')} 元；"
        f"字段识别警告 {len(coverage.get('validation_warnings', []))} 项"
    )
    text = f"""# 报销交付清单

项目：{manifest["project_name"]}
报销人：{manifest["reimburser"]}
费用笔数：{len(manifest["entries"])}
报销总额：{total:.2f}

## 给财务

- 财务提交文件夹：{dirs["finance"]}
- 财务提交 zip：{finance_zip}
- 报销明细表：{path_line(reimbursement)}
- 费用报销单：{form_status}，{path_line(expense_form)}
- 费用截图：{len(copied_screenshots)} 个文件
- 发票：{invoice_status}
- 发票覆盖：{coverage_text}
- 发票缺口清单：{path_line(coverage.get("markdown_report"))}

## 打印

- 一键打印 PDF：{path_line(combined_print_pdf)}，共 {combined_pages} 页
- 费用报销单 PDF：{path_line(expense_form_pdf)}，共 {expense_form_pages} 页，已置于一键打印文件首页
- 费用截图 PDF：{path_line(screenshot_pdf)}，三张一页，共 {screenshot_pages} 页
- 发票 PDF：{path_line(invoice_pdf)}，一张一页，共 {invoice_pages} 页
- 打印排版 Excel：{path_line(print_workbook)}

## 分类合计

{category_text}
"""
    root_checklist = dirs["root"] / "交付清单.md"
    finance_checklist = dirs["finance"] / "交付清单.md"
    for output in [root_checklist, finance_checklist]:
        output.write_text(text, encoding="utf-8")
    return root_checklist


def write_audit_summary(manifest, dirs, combined_print_pdf):
    total = sum((money(entry["amount"]) for entry in manifest["entries"]), Decimal("0.00"))
    coverage = manifest.get("invoice_coverage") or {}
    approval = manifest.get("approval_metadata") or {}
    combined_pages = pdf_page_count(combined_print_pdf) if combined_print_pdf and combined_print_pdf.exists() else 0
    categories = {category: Decimal("0.00") for category in categories_for(manifest)}
    for entry in manifest["entries"]:
        categories.setdefault(entry["expense_type"], Decimal("0.00"))
        categories[entry["expense_type"]] += money(entry["amount"])
    category_text = "、".join(
        f"{category} {amount:.2f} 元"
        for category, amount in categories.items()
        if amount != Decimal("0.00")
    ) or "无"
    unresolved = [
        item
        for item in manifest.get("expense_exceptions", []) or []
        if item.get("severity") == "blocking" and not item.get("resolved")
    ]
    voucher_ids = [entry.get("voucher_id", "") for entry in manifest["entries"]]
    voucher_range = f"{voucher_ids[0]} 至 {voucher_ids[-1]}" if voucher_ids else "无"
    text = f"""# 报销审计摘要

项目：{manifest['project_name']}
报销人：{manifest['reimburser']}
钉钉审批编号：{approval.get('dingding_number', '')}
报销类型：{approval.get('reimbursement_type', '')}
费用笔数：{len(manifest['entries'])}
凭证编号：{voucher_range}
报销总额：{total:.2f} 元
钉钉审批金额：{money(approval.get('approved_amount', 0)):.2f} 元
分类合计：{category_text}
有效发票金额：{coverage.get('recognized_amount', '0.00')} 元
发票缺口：{coverage.get('missing_amount', '0.00')} 元
替代发票：{coverage.get('replacement_invoice_count', 0)} 张，{coverage.get('replacement_invoice_amount', '0.00')} 元
发票字段识别警告：{len(coverage.get('validation_warnings', []))} 项
未解决异常：{len(unresolved)} 项
打印总页数：{combined_pages} 页
自动校验：待执行

说明：系统执行金额、分类、发票覆盖、重复文件、表单字段和打印页数校验。发票真伪仍以公司财务或税务查验平台结果为准。
"""
    root_output = dirs["root"] / "报销审计摘要.md"
    finance_output = dirs["finance"] / "报销审计摘要.md"
    for output in [root_output, finance_output]:
        output.write_text(text, encoding="utf-8")
    return root_output


def write_summary(manifest, dirs, generated, copied_screenshots, copied_invoices, copied_supporting_documents, finance_zip, delivery_checklist, audit_summary):
    category_totals = {category: Decimal("0.00") for category in categories_for(manifest)}
    for entry in manifest["entries"]:
        category = entry["expense_type"]
        category_totals.setdefault(category, Decimal("0.00"))
        category_totals[category] += money(entry["amount"])
    coverage = manifest.get("invoice_coverage") or {}
    summary = {
        "project_name": manifest["project_name"],
        "reimburser": manifest["reimburser"],
        "expense_count": len(manifest["entries"]),
        "screenshot_count": len(copied_screenshots),
        "invoice_count": len(copied_invoices),
        "supporting_document_count": len(copied_supporting_documents),
        "invoice_coverage_status": coverage.get("status", ""),
        "invoice_required_amount": coverage.get("required_amount", ""),
        "invoice_recognized_amount": coverage.get("recognized_amount", ""),
        "invoice_missing_amount": coverage.get("missing_amount", ""),
        "replacement_invoice_count": coverage.get("replacement_invoice_count", 0),
        "replacement_invoice_amount": coverage.get("replacement_invoice_amount", "0.00"),
        "invoice_coverage_report": coverage.get("report", ""),
        "invoice_gap_report": coverage.get("markdown_report", ""),
        "total_amount": str(sum((money(e["amount"]) for e in manifest["entries"]), Decimal("0.00"))),
        "category_totals": {k: str(v.quantize(Decimal("0.01"))) for k, v in category_totals.items()},
        "generated_files": [str(p) for p in generated if p],
        "finance_folder": str(dirs["finance"]),
        "finance_zip": str(finance_zip),
        "delivery_checklist": str(delivery_checklist),
        "audit_summary": str(audit_summary),
    }
    output = dirs["root"] / "pack_summary.json"
    output.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return output


def main():
    args = parse_args()
    manifest = load_manifest(args.manifest)
    validate_manifest(manifest)
    root = Path(args.out).expanduser().resolve()
    dirs = ensure_dirs(root)
    prepare_manifest_images(manifest, dirs)

    reimbursement = build_reimbursement_workbook(manifest, dirs)
    expense_form = build_expense_form(manifest, dirs)
    expense_form_pdf = build_expense_form_pdf(manifest, dirs, expense_form)
    if expense_form_pdf:
        actual_form_pages = pdf_page_count(expense_form_pdf)
        expected_form_pages = manifest_context(manifest)["expense_form_print_pages"]
        if actual_form_pages != expected_form_pages:
            manifest["_expense_form_print_pages"] = actual_form_pages
            expense_form = build_expense_form(manifest, dirs)
            expense_form_pdf = build_expense_form_pdf(manifest, dirs, expense_form)
            if pdf_page_count(expense_form_pdf) != actual_form_pages:
                raise SystemExit("费用报销单页数在更新总页数后发生变化，请检查模板打印区域。")
    screenshot_pdf = build_screenshot_pdf(manifest, dirs)
    invoice_pdf = build_invoice_pdf(manifest, dirs)
    combined_print_pdf = build_combined_print_pdf(manifest, dirs, expense_form_pdf, screenshot_pdf, invoice_pdf)
    print_workbook = build_print_workbook(manifest, dirs)
    generated = [reimbursement, expense_form, expense_form_pdf, screenshot_pdf, invoice_pdf, combined_print_pdf, print_workbook]
    copied_screenshots, copied_invoices, copied_supporting_documents = copy_sources_and_outputs(manifest, dirs, generated)
    finance_zip_path = expected_finance_zip_path(manifest, dirs)
    delivery_checklist = write_delivery_checklist(
        manifest,
        dirs,
        reimbursement,
        expense_form,
        expense_form_pdf,
        screenshot_pdf,
        invoice_pdf,
        combined_print_pdf,
        print_workbook,
        copied_screenshots,
        copied_invoices,
        copied_supporting_documents,
        finance_zip_path,
    )
    audit_summary = write_audit_summary(manifest, dirs, combined_print_pdf)
    finance_zip = build_finance_zip(manifest, dirs)
    summary = write_summary(
        manifest,
        dirs,
        generated,
        copied_screenshots,
        copied_invoices,
        copied_supporting_documents,
        finance_zip,
        delivery_checklist,
        audit_summary,
    )

    print(json.dumps({
        "output_folder": str(root),
        "finance_folder": str(dirs["finance"]),
        "finance_zip": str(finance_zip),
        "delivery_checklist": str(delivery_checklist),
        "summary": str(summary),
        "generated_files": [str(p) for p in generated if p],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    from runtime_utils import configure_utf8_stdio

    configure_utf8_stdio()
    main()
