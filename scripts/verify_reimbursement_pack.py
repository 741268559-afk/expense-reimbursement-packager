#!/usr/bin/env python3
import argparse
import json
import math
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

from build_reimbursement_pack import chinese_rmb_upper


def parse_args():
    parser = argparse.ArgumentParser(description="Verify a generated reimbursement pack.")
    parser.add_argument("--pack", required=True, help="Generated pack folder.")
    parser.add_argument("--manifest", help="Optional manifest path. Defaults to draft_manifest.json or finance manifest.json.")
    parser.add_argument("--out", help="Optional verification JSON output path. Defaults to <pack>/verification.json.")
    return parser.parse_args()


def money(value):
    return Decimal(str(value)).quantize(Decimal("0.01"))


def decimal_value(value):
    if isinstance(value, bool) or value in (None, ""):
        raise InvalidOperation
    parsed = Decimal(str(value).strip())
    if not parsed.is_finite():
        raise InvalidOperation
    return parsed


def safe_name(text, fallback="item"):
    text = str(text or "").strip()
    text = re.sub(r"[\\/:*?\"<>|]+", "_", text)
    text = re.sub(r"\s+", "", text)
    text = text.strip("._-")
    return text or fallback


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def check(checks, name, passed, details=""):
    checks.append({
        "name": name,
        "status": "passed" if passed else "failed",
        "details": details,
    })


def find_generated(summary, needle):
    for value in summary.get("generated_files", []):
        path = Path(value)
        if needle in path.name:
            return path
    return None


def find_default_manifest(pack, manifest_arg):
    if manifest_arg:
        return Path(manifest_arg).expanduser().resolve()
    candidates = [
        pack / "draft_manifest.json",
        pack / "财务提交文件夹" / "manifest.json",
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def manifest_totals(manifest):
    total = Decimal("0.00")
    categories = {category: Decimal("0.00") for category in manifest.get("categories", [])}
    category_counts = {category: 0 for category in manifest.get("categories", [])}
    screenshot_count = 0
    invoices = []
    for invoice in manifest.get("invoices", []) or []:
        if invoice:
            invoices.append(str(invoice))
    for entry in manifest.get("entries", []) or []:
        try:
            amount = money(entry.get("amount"))
        except (InvalidOperation, ValueError, TypeError):
            amount = Decimal("0.00")
        total += amount
        category = entry.get("expense_type", "")
        categories[category] = categories.get(category, Decimal("0.00")) + amount
        category_counts[category] = category_counts.get(category, 0) + 1
        if entry.get("screenshot"):
            screenshot_count += 1
        for invoice in entry.get("invoices", []) or []:
            if invoice:
                invoices.append(str(invoice))
    unique_invoices = []
    seen = set()
    for invoice in invoices:
        if invoice in seen:
            continue
        seen.add(invoice)
        unique_invoices.append(invoice)
    supporting_documents = []
    seen_supporting = set()
    for value in manifest.get("supporting_documents", []) or []:
        if not value or str(value) in seen_supporting:
            continue
        seen_supporting.add(str(value))
        supporting_documents.append(str(value))
    return {
        "expense_count": len(manifest.get("entries", []) or []),
        "screenshot_count": screenshot_count,
        "invoice_count": len(unique_invoices),
        "supporting_document_count": len(supporting_documents),
        "total_amount": total,
        "category_totals": categories,
        "category_counts": category_counts,
    }


def manifest_integrity_issues(manifest, manifest_path):
    issues = []
    entries = manifest.get("entries", []) or []
    if not entries:
        issues.append("manifest has no expense entries")
    if not str(manifest.get("reimburser", "")).strip():
        issues.append("manifest.reimburser is blank")

    base_dir = manifest_path.parent if manifest_path else Path.cwd()
    screenshot_rows = {}
    voucher_ids = []
    for index, entry in enumerate(entries, start=1):
        for field in ["date", "purpose", "expense_type", "amount", "screenshot"]:
            if entry.get(field) in (None, ""):
                issues.append(f"entry {index}: {field} is blank")

        date_text = str(entry.get("date", "")).strip()
        try:
            datetime.strptime(date_text, "%Y-%m-%d")
        except ValueError:
            issues.append(f"entry {index}: invalid date {date_text!r}")

        try:
            raw_amount = decimal_value(entry.get("amount"))
            if raw_amount <= 0:
                issues.append(f"entry {index}: amount must be positive, got {raw_amount}")
            if raw_amount != raw_amount.quantize(Decimal("0.01")):
                issues.append(f"entry {index}: amount has more than two decimal places, got {raw_amount}")
            ocr_text = str(entry.get("ocr_text", ""))
            if ocr_text:
                normalized_ocr = ocr_text.replace(",", "").replace("，", "").replace(" ", "")
                expected_text = f"{raw_amount.quantize(Decimal('0.01')):.2f}"
                if expected_text not in normalized_ocr:
                    issues.append(f"entry {index}: amount {expected_text} not found in OCR text")
        except (InvalidOperation, ValueError, TypeError):
            issues.append(f"entry {index}: invalid amount {entry.get('amount')!r}")

        if "NEEDS_REVIEW" in str(entry.get("notes", "")):
            issues.append(f"entry {index}: still marked NEEDS_REVIEW")

        voucher_id = str(entry.get("voucher_id", "")).strip()
        if manifest.get("workflow_version", 1) >= 2 and not re.fullmatch(r"FY\d{3,}", voucher_id):
            issues.append(f"entry {index}: invalid voucher_id {voucher_id!r}")
        if voucher_id:
            voucher_ids.append(voucher_id)

        screenshot = entry.get("screenshot")
        if screenshot:
            screenshot_path = Path(str(screenshot)).expanduser()
            if not screenshot_path.is_absolute():
                screenshot_path = base_dir / screenshot_path
            screenshot_key = str(screenshot_path.resolve())
            screenshot_rows.setdefault(screenshot_key, []).append(index)

    for screenshot, rows in screenshot_rows.items():
        if len(rows) > 1:
            issues.append(f"duplicate screenshot used by entries {rows}: {screenshot}")
    if len(voucher_ids) != len(set(voucher_ids)):
        issues.append("voucher_id values are not unique")
    unresolved_exceptions = [
        item
        for item in manifest.get("expense_exceptions", []) or []
        if item.get("severity") == "blocking" and not item.get("resolved")
    ]
    if unresolved_exceptions:
        issues.append(f"{len(unresolved_exceptions)} blocking expense exceptions remain unresolved")
    approval_requirement = manifest.get("approval_requirement") or {}
    if approval_requirement.get("policy", "none") == "required":
        approval = manifest.get("approval_metadata") or {}
        for field in approval_requirement.get("required_fields", []) or []:
            if approval.get(field) in (None, ""):
                issues.append(f"approval_metadata.{field} is blank")
        total = sum((money(entry.get("amount", 0)) for entry in entries), Decimal("0.00"))
        try:
            approved_amount = money(approval.get("approved_amount"))
        except (InvalidOperation, ValueError, TypeError):
            approved_amount = None
        if approved_amount != total:
            issues.append(f"DingTalk approved amount {approved_amount} != reimbursement total {total}")
        confirmation = manifest.get("expense_confirmation") or {}
        try:
            confirmed_amount = money(confirmation.get("confirmed_amount"))
        except (InvalidOperation, ValueError, TypeError):
            confirmed_amount = None
        if confirmation.get("status") != "confirmed_by_dingtalk" or confirmed_amount != total:
            issues.append("expense confirmation does not match the DingTalk-approved total")
    for index, value in enumerate(manifest.get("supporting_documents", []) or [], start=1):
        path = Path(str(value)).expanduser()
        if not path.is_absolute():
            path = base_dir / path
        if not path.exists():
            issues.append(f"supporting document {index}: file not found: {path}")
    return issues


def normalize_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value or "").strip()


def manifest_detail_rows(manifest):
    entries = sorted(
        manifest.get("entries", []) or [],
        key=lambda entry: (
            entry.get("date") or "9999-12-31",
            entry.get("time") or "23:59:59",
            entry.get("screenshot") or entry.get("source") or "",
        ),
    )
    reimburser = str(manifest.get("reimburser", "")).strip()
    rows = []
    for index, entry in enumerate(entries, start=1):
        try:
            amount = money(entry.get("amount"))
        except (InvalidOperation, ValueError, TypeError):
            amount = Decimal("0.00")
        invoice_names = [Path(str(value)).name for value in (entry.get("invoices", []) or []) if value]
        rows.append({
            "sequence": entry.get("voucher_id") or f"FY{index:03d}",
            "date": normalize_date(entry.get("date")),
            "purpose": str(entry.get("purpose", "")).strip(),
            "reimburser": reimburser,
            "expense_type": str(entry.get("expense_type", "")).strip(),
            "amount": amount,
            "invoice": "；".join(invoice_names),
        })
    return rows


def workbook_totals(path):
    wb = load_workbook(path, data_only=True)
    ws = wb["报销明细"] if "报销明细" in wb.sheetnames else wb.active
    rows = []
    headers = []
    header_row = None
    value_issues = []
    for row in range(1, ws.max_row + 1):
        values = [ws.cell(row, col).value for col in range(1, 9)]
        if values[:8] in [
            ["凭证编号", "日期", "用途", "报销人", "费用类型", "金额(元)", "凭证截图", "发票"],
            ["序号", "日期", "用途", "报销人", "费用类型", "金额(元)", "凭证截图", "发票"],
        ]:
            headers = values
            header_row = row
        first_value = ws.cell(row, 1).value
        if type(first_value) is int or re.fullmatch(r"FY\d{3,}", str(first_value or "")):
            try:
                amount = money(ws.cell(row, 6).value)
            except (InvalidOperation, ValueError, TypeError):
                amount = Decimal("0.00")
                value_issues.append(f"row {row}: invalid amount {ws.cell(row, 6).value!r}")
            rows.append({
                "sequence": ws.cell(row, 1).value,
                "date": normalize_date(ws.cell(row, 2).value),
                "purpose": str(ws.cell(row, 3).value or "").strip(),
                "reimburser": str(ws.cell(row, 4).value or "").strip(),
                "expense_type": str(ws.cell(row, 5).value or "").strip(),
                "amount": amount,
                "invoice": str(ws.cell(row, 8).value or "").strip(),
                "sheet_row": row,
            })
    total = sum((row["amount"] for row in rows), Decimal("0.00"))
    categories = {}
    for row in rows:
        category = row["expense_type"]
        categories[category] = categories.get(category, Decimal("0.00")) + row["amount"]

    summary_categories = {}
    summary_total = None
    summary_count = None
    if header_row:
        for row in range(3, header_row):
            label = str(ws.cell(row, 1).value or "").strip()
            if not label:
                continue
            try:
                amount = money(ws.cell(row, 5).value)
            except (InvalidOperation, ValueError, TypeError):
                amount = Decimal("0.00")
                value_issues.append(f"summary row {row}: invalid amount {ws.cell(row, 5).value!r}")
            count_match = re.search(r"(\d+)\s*笔", str(ws.cell(row, 7).value or ""))
            count = int(count_match.group(1)) if count_match else None
            if label.replace(" ", "") == "总计":
                summary_total = amount
                summary_count = count
            else:
                summary_categories[label] = {"amount": amount, "count": count}

    footer_total = None
    for row in range(1, ws.max_row + 1):
        if str(ws.cell(row, 1).value or "").strip() == "合计":
            try:
                footer_total = money(ws.cell(row, 6).value)
            except (InvalidOperation, ValueError, TypeError):
                value_issues.append(f"footer row {row}: invalid amount {ws.cell(row, 6).value!r}")
            break

    image_rows = []
    for image in getattr(ws, "_images", []):
        marker = getattr(getattr(image, "anchor", None), "_from", None)
        if marker is not None:
            image_rows.append(marker.row + 1)

    formula_errors = []
    formula_wb = load_workbook(path, data_only=False)
    error_pattern = re.compile(r"#(?:REF!|DIV/0!|VALUE!|NAME\?|N/A|NUM!|NULL!|SPILL!|CALC!)")
    for formula_ws in formula_wb.worksheets:
        for formula_row in formula_ws.iter_rows():
            for cell in formula_row:
                if cell.data_type == "e" or (isinstance(cell.value, str) and error_pattern.search(cell.value)):
                    formula_errors.append(f"{formula_ws.title}!{cell.coordinate}={cell.value}")
    return {
        "title": ws["A1"].value,
        "headers": headers,
        "header_row": header_row,
        "rows": rows,
        "expense_count": len(rows),
        "total_amount": total,
        "category_totals": categories,
        "image_count": len(getattr(ws, "_images", [])),
        "image_rows": sorted(image_rows),
        "summary_categories": summary_categories,
        "summary_total": summary_total,
        "summary_count": summary_count,
        "footer_total": footer_total,
        "value_issues": value_issues,
        "formula_errors": formula_errors,
    }


def detail_row_mismatches(actual_rows, expected_rows):
    mismatches = []
    if len(actual_rows) != len(expected_rows):
        mismatches.append(f"row count {len(actual_rows)} != {len(expected_rows)}")
    fields = ["sequence", "date", "purpose", "reimburser", "expense_type", "amount", "invoice"]
    for index, (actual, expected) in enumerate(zip(actual_rows, expected_rows), start=1):
        for field in fields:
            if field == "amount":
                matches = actual[field] == expected[field]
            else:
                matches = actual[field] == expected[field]
            if not matches:
                mismatches.append(f"row {index} {field}: {actual[field]!r} != {expected[field]!r}")
    return mismatches


def summary_row_mismatches(actual, expected):
    mismatches = []
    expected_categories = set(expected["category_totals"])
    actual_categories = set(actual)
    if actual_categories != expected_categories:
        mismatches.append(f"category labels {sorted(actual_categories)} != {sorted(expected_categories)}")
    for category in sorted(expected_categories):
        observed = actual.get(category, {})
        if observed.get("amount") != expected["category_totals"][category]:
            mismatches.append(
                f"{category} amount: {observed.get('amount')!r} != {expected['category_totals'][category]!r}"
            )
        if observed.get("count") != expected["category_counts"][category]:
            mismatches.append(
                f"{category} count: {observed.get('count')!r} != {expected['category_counts'][category]!r}"
            )
    return mismatches


def pdf_page_count(path):
    return len(PdfReader(str(path)).pages)


def form_placeholders(path):
    wb = load_workbook(path, data_only=False)
    leftovers = []
    pattern = re.compile(r"\{\{[^{}]+?\}\}")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and pattern.search(cell.value):
                    leftovers.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    return leftovers


def writable_cell_ref(ws, cell_ref):
    for cell_range in ws.merged_cells.ranges:
        if cell_ref in cell_range:
            return f"{get_column_letter(cell_range.min_col)}{cell_range.min_row}"
    return cell_ref


def manifest_context(manifest, expected, total_print_pages=None):
    dates = [entry.get("date") for entry in manifest.get("entries", []) or [] if entry.get("date")]
    category_counts = {category: 0 for category in expected["category_totals"]}
    for entry in manifest.get("entries", []) or []:
        category = entry.get("expense_type", "")
        category_counts[category] = category_counts.get(category, 0) + 1

    approval = manifest.get("approval_metadata") or {}
    context = {
        "project_name": manifest.get("project_name", ""),
        "reimburser": manifest.get("reimburser", ""),
        "dingding_number": approval.get("dingding_number", ""),
        "reimbursement_type": approval.get("reimbursement_type", ""),
        "payee": approval.get("payee", ""),
        "invoice_entity": approval.get("invoice_entity", ""),
        "approval_company": approval.get("company", ""),
        "reimbursement_project": approval.get("reimbursement_project", ""),
        "contract_number": approval.get("contract_number", ""),
        "reimbursement_reason": approval.get("reimbursement_reason", ""),
        "business_line": approval.get("business_line", ""),
        "approved_amount": f"{money(approval.get('approved_amount')):.2f}" if approval.get("approved_amount") not in (None, "") else "",
        "expense_count": expected["expense_count"],
        "total_amount": f"{expected['total_amount']:.2f}",
        "date_start": min(dates) if dates else "",
        "date_end": max(dates) if dates else "",
        "date_range": f"{min(dates)}至{max(dates)}" if dates else "",
        "screenshot_print_pages": math.ceil(expected["screenshot_count"] / 3) if expected["screenshot_count"] else 0,
        "invoice_print_pages": expected["invoice_count"] + expected["supporting_document_count"],
        "expense_form_print_pages": 1 if (manifest.get("reimbursement_form") or {}).get("template") else 0,
        "total_print_pages": total_print_pages if total_print_pages is not None else (
            (1 if (manifest.get("reimbursement_form") or {}).get("template") else 0)
            +
            (math.ceil(expected["screenshot_count"] / 3) if expected["screenshot_count"] else 0)
            + expected["invoice_count"]
            + expected["supporting_document_count"]
        ),
    }
    for category, amount in expected["category_totals"].items():
        key = safe_name(category)
        context[f"{key}_amount"] = f"{amount:.2f}"
        context[f"{key}_count"] = category_counts.get(category, 0)
    return context


def render_form_value(value, context):
    if isinstance(value, str):
        def replace(match):
            key = match.group(1).strip()
            return str(context.get(key, match.group(0)))
        rendered = re.sub(r"\{\{\s*([^{}]+?)\s*\}\}", replace, value)
        if re.fullmatch(r"-?\d+(?:\.\d+)?", rendered):
            integer_part = rendered.lstrip("-").split(".", 1)[0]
            if len(integer_part) > 15:
                return rendered
            return money(rendered)
        return rendered
    if isinstance(value, (int, float, Decimal)):
        return money(value)
    return value


def values_match(actual, expected):
    if expected in (None, ""):
        return actual in (None, "")
    if isinstance(expected, Decimal):
        try:
            return money(actual) == expected
        except Exception:
            return False
    return str(actual or "").strip() == str(expected).strip()


def form_mapping_mismatches(path, form, context, manifest=None):
    wb = load_workbook(path, data_only=True)
    sheet_name = form.get("sheet")
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
    mismatches = []
    for cell_ref, configured_value in (form.get("cells") or {}).items():
        actual_ref = writable_cell_ref(ws, cell_ref)
        expected = render_form_value(configured_value, context)
        actual = ws[actual_ref].value
        if not values_match(actual, expected):
            mismatches.append(f"{ws.title}!{actual_ref}: {actual!r} != {expected!r}")
    category_rows = form.get("category_rows") or []
    if category_rows and manifest is not None:
        categories = list(manifest.get("categories") or [])
        for entry in manifest.get("entries", []) or []:
            category = entry.get("expense_type", "")
            if category and category not in categories:
                categories.append(category)
        expected_categories = []
        for category in categories:
            amount = money(context.get(f"{safe_name(category)}_amount", "0.00"))
            if amount != Decimal("0.00"):
                expected_categories.append((category, amount))
        if len(expected_categories) > len(category_rows):
            mismatches.append(f"category rows: {len(expected_categories)} nonzero categories > {len(category_rows)} available rows")
        for index, row_config in enumerate(category_rows):
            label_ref = writable_cell_ref(ws, row_config.get("label_cell", ""))
            amount_ref = writable_cell_ref(ws, row_config.get("amount_cell", ""))
            expected_label, expected_amount = expected_categories[index] if index < len(expected_categories) else ("", "")
            if not values_match(ws[label_ref].value, expected_label):
                mismatches.append(f"{ws.title}!{label_ref}: {ws[label_ref].value!r} != {expected_label!r}")
            if not values_match(ws[amount_ref].value, expected_amount):
                mismatches.append(f"{ws.title}!{amount_ref}: {ws[amount_ref].value!r} != {expected_amount!r}")
    return mismatches


def form_formula_cache_mismatches(path, expected_total):
    formula_wb = load_workbook(path, data_only=False)
    cached_wb = load_workbook(path, data_only=True)
    sum_pattern = re.compile(r"^=?SUM\(\s*([A-Z]{1,3}\d+):([A-Z]{1,3}\d+)\s*\)$", re.IGNORECASE)
    mismatches = []
    checked = 0
    for formula_ws, cached_ws in zip(formula_wb.worksheets, cached_wb.worksheets):
        for row in formula_ws.iter_rows():
            for cell in row:
                if cell.data_type != "f" or not isinstance(cell.value, str):
                    continue
                formula = cell.value.strip()
                match = sum_pattern.fullmatch(formula)
                expected_value = None
                if match:
                    expected_value = Decimal("0.00")
                    for range_row in formula_ws[f"{match.group(1)}:{match.group(2)}"]:
                        for source_cell in range_row:
                            if isinstance(source_cell.value, (int, float, Decimal)) and not isinstance(source_cell.value, bool):
                                expected_value += money(source_cell.value)
                    expected_value = money(expected_value)
                elif "[DBNUM2]" in formula.upper():
                    expected_value = chinese_rmb_upper(expected_total)
                if expected_value is None:
                    continue
                checked += 1
                actual = cached_ws[cell.coordinate].value
                if not values_match(actual, expected_value):
                    mismatches.append(f"{formula_ws.title}!{cell.coordinate}: {actual!r} != {expected_value!r}")
    return checked, mismatches


def zip_report(path):
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
    bad = [
        name for name in names
        if "/._" in name or name.startswith("._") or name.endswith(".DS_Store") or "__MACOSX" in name or "__pycache__" in name or name.endswith(".pyc")
    ]
    return {
        "names": names,
        "bad": bad,
        "has_tables_dir": any(name.endswith("/01_报销表/") for name in names),
        "has_screenshots_dir": any(name.endswith("/02_费用截图/") for name in names),
        "has_invoices_dir": any(name.endswith("/03_发票/") for name in names),
        "has_print_dir": any(name.endswith("/04_打印文件/") for name in names),
        "has_delivery_checklist": any(name.endswith("/交付清单.md") for name in names),
        "has_expense_confirmation": any(name.endswith("/报销金额确认单.md") for name in names),
        "has_expense_exceptions": any(name.endswith("/报销异常清单.md") for name in names),
        "has_approval_gap": any(name.endswith("/钉钉信息缺口.md") for name in names),
        "has_audit_summary": any(name.endswith("/报销审计摘要.md") for name in names),
        "has_invoice_coverage_report": any(name.endswith("/invoice_coverage.json") for name in names),
        "has_invoice_gap_report": any(name.endswith("/发票缺口清单.md") for name in names),
        "screenshot_files": sum("/02_费用截图/" in name and not name.endswith("/") for name in names),
        "screenshot_names": [name for name in names if "/02_费用截图/" in name and not name.endswith("/")],
        "invoice_files": sum("/03_发票/" in name and not name.endswith("/") for name in names),
        "replacement_invoice_files": sum("/03_发票/" in name and "替代发票" in name and not name.endswith("/") for name in names),
        "supporting_document_files": sum("/03_发票/" in name and "发票附件_行程单" in name and not name.endswith("/") for name in names),
        "table_files": sum("/01_报销表/" in name and not name.endswith("/") for name in names),
        "print_files": sum("/04_打印文件/" in name and not name.endswith("/") for name in names),
    }


def decimal_dict(text_dict):
    parsed = {}
    issues = []
    for key, value in (text_dict or {}).items():
        try:
            parsed[key] = money(value)
        except (InvalidOperation, ValueError, TypeError):
            issues.append(f"{key}: invalid amount {value!r}")
    return parsed, issues


def main():
    args = parse_args()
    pack = Path(args.pack).expanduser().resolve()
    checks = []
    summary_path = pack / "pack_summary.json"
    check(checks, "pack folder exists", pack.exists(), str(pack))
    check(checks, "pack_summary.json exists", summary_path.exists(), str(summary_path))
    if not summary_path.exists():
        result = {"status": "failed", "pack": str(pack), "checks": checks}
        print(json.dumps(result, ensure_ascii=False, indent=2))
        raise SystemExit(1)

    summary = read_json(summary_path)
    manifest_path = find_default_manifest(pack, args.manifest)
    check(checks, "manifest exists", bool(manifest_path and manifest_path.exists()), str(manifest_path or ""))
    manifest = read_json(manifest_path) if manifest_path and manifest_path.exists() else {"entries": []}
    integrity_issues = manifest_integrity_issues(manifest, manifest_path)
    check(
        checks,
        "manifest amounts and source references are valid",
        not integrity_issues,
        "; ".join(integrity_issues[:10]),
    )
    expected = manifest_totals(manifest)
    expected_rows = manifest_detail_rows(manifest)

    try:
        summary_total = money(summary.get("total_amount"))
        summary_total_issue = ""
    except (InvalidOperation, ValueError, TypeError):
        summary_total = Decimal("0.00")
        summary_total_issue = f"invalid amount {summary.get('total_amount')!r}"
    check(checks, "summary total is a valid amount", not summary_total_issue, summary_total_issue)
    check(checks, "summary total matches manifest", summary_total == expected["total_amount"], f"{summary_total} vs {expected['total_amount']}")
    check(checks, "summary expense count matches manifest", summary.get("expense_count") == expected["expense_count"], f"{summary.get('expense_count')} vs {expected['expense_count']}")
    check(checks, "summary screenshot count matches manifest", summary.get("screenshot_count") == expected["screenshot_count"], f"{summary.get('screenshot_count')} vs {expected['screenshot_count']}")
    check(checks, "summary invoice count matches manifest", summary.get("invoice_count") == expected["invoice_count"], f"{summary.get('invoice_count')} vs {expected['invoice_count']}")
    check(
        checks,
        "summary supporting document count matches manifest",
        summary.get("supporting_document_count", 0) == expected["supporting_document_count"],
        f"{summary.get('supporting_document_count', 0)} vs {expected['supporting_document_count']}",
    )

    summary_categories, summary_category_issues = decimal_dict(summary.get("category_totals", {}))
    check(checks, "summary category amounts are valid", not summary_category_issues, "; ".join(summary_category_issues[:10]))
    check(
        checks,
        "summary category labels match manifest",
        set(summary_categories) == set(expected["category_totals"]),
        f"{sorted(summary_categories)} vs {sorted(expected['category_totals'])}",
    )
    for category, amount in expected["category_totals"].items():
        check(checks, f"summary category {category}", summary_categories.get(category, Decimal("0.00")) == amount, f"{summary_categories.get(category, Decimal('0.00'))} vs {amount}")

    requirement = manifest.get("invoice_requirement") or {}
    coverage = manifest.get("invoice_coverage") or {}
    invoice_policy = requirement.get("policy") or "full_amount"
    coverage_required = invoice_policy == "full_amount"
    coverage_amounts = {}
    replacement_invoice_count = 0
    if coverage_required:
        coverage_issues = []
        for key in ["expense_total", "required_amount", "recognized_amount", "missing_amount", "excess_amount", "replacement_invoice_amount"]:
            try:
                coverage_amounts[key] = money(coverage.get(key))
            except (InvalidOperation, ValueError, TypeError):
                coverage_issues.append(f"{key}: invalid amount {coverage.get(key)!r}")
        check(checks, "invoice coverage amounts are valid", not coverage_issues, "; ".join(coverage_issues[:10]))
        check(checks, "invoice coverage status is covered", coverage.get("status") == "covered", str(coverage.get("status", "")))
        check(checks, "invoice required amount equals reimbursement total", coverage_amounts.get("required_amount") == expected["total_amount"], f"{coverage_amounts.get('required_amount')} vs {expected['total_amount']}")
        check(checks, "invoice coverage expense total matches manifest", coverage_amounts.get("expense_total") == expected["total_amount"], f"{coverage_amounts.get('expense_total')} vs {expected['total_amount']}")
        check(checks, "invoice missing amount is zero", coverage_amounts.get("missing_amount") == Decimal("0.00"), str(coverage_amounts.get("missing_amount")))
        recognized_covers_required = (
            not coverage_issues
            and coverage_amounts.get("recognized_amount", Decimal("0.00"))
            >= coverage_amounts.get("required_amount", Decimal("0.00"))
        )
        check(checks, "recognized invoices cover required amount", recognized_covers_required, f"{coverage_amounts.get('recognized_amount')} vs {coverage_amounts.get('required_amount')}")

        invoice_items = manifest.get("invoice_items", []) or []
        unresolved_items = [item for item in invoice_items if item.get("status") != "recognized"]
        check(checks, "all invoice files have recognized non-duplicate totals", not unresolved_items, "; ".join(f"{Path(item.get('path', '')).name}: {item.get('status')}" for item in unresolved_items[:10]))
        traceability_issues = []
        for item in invoice_items:
            for field in ["sha256", "invoice_number_status", "invoice_date_status", "entity_check", "authenticity_check", "voucher_ids"]:
                if field not in item:
                    traceability_issues.append(f"{Path(item.get('path', '')).name}: missing {field}")
        check(checks, "invoice traceability fields are recorded", not traceability_issues, "; ".join(traceability_issues[:10]))
        check(
            checks,
            "invoice authenticity boundary is explicit",
            all(item.get("authenticity_check") == "not_performed" for item in invoice_items),
            "Tax-platform authenticity verification is external.",
        )
        check(checks, "invoice coverage item count matches manifest", len(invoice_items) == expected["invoice_count"], f"{len(invoice_items)} vs {expected['invoice_count']}")
        replacement_items = [item for item in invoice_items if item.get("role") == "replacement" and item.get("status") == "recognized"]
        replacement_invoice_count = len(replacement_items)
        replacement_total = sum((money(item.get("amount")) for item in replacement_items), Decimal("0.00"))
        check(checks, "replacement invoice count matches coverage", coverage.get("replacement_invoice_count") == replacement_invoice_count, f"{coverage.get('replacement_invoice_count')} vs {replacement_invoice_count}")
        check(checks, "replacement invoice amount matches coverage", coverage_amounts.get("replacement_invoice_amount") == replacement_total, f"{coverage_amounts.get('replacement_invoice_amount')} vs {replacement_total}")

        coverage_report = Path(coverage.get("report", "")) if coverage.get("report") else None
        gap_report = Path(coverage.get("markdown_report", "")) if coverage.get("markdown_report") else None
        check(checks, "invoice coverage JSON exists", bool(coverage_report and coverage_report.exists()), str(coverage_report or ""))
        check(checks, "invoice gap report exists", bool(gap_report and gap_report.exists()), str(gap_report or ""))
        check(checks, "summary invoice coverage status", summary.get("invoice_coverage_status") == "covered", str(summary.get("invoice_coverage_status", "")))
        check(checks, "summary invoice required amount", money(summary.get("invoice_required_amount", 0)) == coverage_amounts.get("required_amount"), f"{summary.get('invoice_required_amount')} vs {coverage.get('required_amount')}")
        check(checks, "summary invoice recognized amount", money(summary.get("invoice_recognized_amount", 0)) == coverage_amounts.get("recognized_amount"), f"{summary.get('invoice_recognized_amount')} vs {coverage.get('recognized_amount')}")
        check(checks, "summary invoice missing amount", money(summary.get("invoice_missing_amount", 0)) == Decimal("0.00"), str(summary.get("invoice_missing_amount")))

    delivery_value = summary.get("delivery_checklist", "")
    delivery_checklist = Path(delivery_value) if delivery_value else None
    check(checks, "delivery checklist exists", bool(delivery_checklist and delivery_checklist.exists()), str(delivery_checklist or ""))
    if delivery_checklist and delivery_checklist.exists():
        checklist_text = delivery_checklist.read_text(encoding="utf-8")
        checklist_ok = "一键打印 PDF" in checklist_text and "财务提交 zip" in checklist_text and str(expected["total_amount"]) in checklist_text
        check(checks, "delivery checklist has finance and print instructions", checklist_ok, delivery_checklist.name)
        category_lines_present = all(
            f"- {category}: {amount:.2f}" in checklist_text
            for category, amount in expected["category_totals"].items()
        )
        check(checks, "delivery checklist category totals match manifest", category_lines_present, delivery_checklist.name)
        if coverage_required:
            invoice_coverage_present = (
                f"应覆盖 {coverage.get('required_amount')} 元" in checklist_text
                and f"缺口 {coverage.get('missing_amount')} 元" in checklist_text
                and f"替代发票 {coverage.get('replacement_invoice_count')} 张" in checklist_text
            )
            check(checks, "delivery checklist has invoice coverage totals", invoice_coverage_present, delivery_checklist.name)

    if manifest.get("workflow_version", 1) >= 2:
        workflow_reports = manifest.get("workflow_reports") or {}
        confirmation_report = Path(workflow_reports.get("expense_confirmation", "")) if workflow_reports.get("expense_confirmation") else None
        exception_report = Path(workflow_reports.get("expense_exceptions", "")) if workflow_reports.get("expense_exceptions") else None
        approval_gap_report = Path(workflow_reports.get("approval_gap", "")) if workflow_reports.get("approval_gap") else None
        check(checks, "expense confirmation report exists", bool(confirmation_report and confirmation_report.exists()), str(confirmation_report or ""))
        if confirmation_report and confirmation_report.exists():
            confirmation_text = confirmation_report.read_text(encoding="utf-8")
            confirmation_ok = (
                f"报销总额：{expected['total_amount']:.2f} 元" in confirmation_text
                and all(str(row["sequence"]) in confirmation_text for row in expected_rows)
            )
            check(checks, "expense confirmation report reconciles total and voucher IDs", confirmation_ok, confirmation_report.name)
        check(checks, "expense exception report exists", bool(exception_report and exception_report.exists()), str(exception_report or ""))
        if exception_report and exception_report.exists():
            check(checks, "expense exception report has no blocking items", "阻断异常：0 项" in exception_report.read_text(encoding="utf-8"), exception_report.name)
        check(checks, "DingTalk gap report exists", bool(approval_gap_report and approval_gap_report.exists()), str(approval_gap_report or ""))
        if approval_gap_report and approval_gap_report.exists():
            check(checks, "DingTalk fields and amount are complete", "当前状态：完整" in approval_gap_report.read_text(encoding="utf-8"), approval_gap_report.name)

    audit_value = summary.get("audit_summary", "")
    audit_summary = Path(audit_value) if audit_value else None
    check(checks, "audit summary exists", bool(audit_summary and audit_summary.exists()), str(audit_summary or ""))
    if audit_summary and audit_summary.exists():
        audit_text = audit_summary.read_text(encoding="utf-8")
        audit_ok = (
            f"报销总额：{expected['total_amount']:.2f} 元" in audit_text
            and ("自动校验：待执行" in audit_text or "自动校验：全部通过" in audit_text)
        )
        check(checks, "audit summary contains reconciled total and verification state", audit_ok, audit_summary.name)

    reimbursement = find_generated(summary, "报销明细表")
    check(checks, "reimbursement workbook exists", bool(reimbursement and reimbursement.exists()), str(reimbursement or ""))
    wb_totals = None
    if reimbursement and reimbursement.exists():
        wb_totals = workbook_totals(reimbursement)
        check(
            checks,
            "reimbursement workbook title",
            wb_totals["title"] == f"{manifest.get('project_name', '')}{manifest.get('reimburser', '')} 报销明细表",
            str(wb_totals["title"]),
        )
        check(checks, "reimbursement workbook header uses purpose", "用途" in wb_totals["headers"], str(wb_totals["headers"]))
        if manifest.get("workflow_version", 1) >= 2:
            check(checks, "reimbursement workbook uses voucher IDs", "凭证编号" in wb_totals["headers"], str(wb_totals["headers"]))
        check(checks, "reimbursement workbook numeric cells are valid", not wb_totals["value_issues"], "; ".join(wb_totals["value_issues"][:10]))
        check(checks, "reimbursement workbook has no formula errors", not wb_totals["formula_errors"], "; ".join(wb_totals["formula_errors"][:10]))
        check(checks, "reimbursement workbook expense count", wb_totals["expense_count"] == expected["expense_count"], f"{wb_totals['expense_count']} vs {expected['expense_count']}")
        row_mismatches = detail_row_mismatches(wb_totals["rows"], expected_rows)
        check(checks, "reimbursement workbook rows match manifest exactly", not row_mismatches, "; ".join(row_mismatches[:10]))
        check(checks, "reimbursement workbook total", wb_totals["total_amount"] == expected["total_amount"], f"{wb_totals['total_amount']} vs {expected['total_amount']}")
        check(checks, "reimbursement workbook footer total", wb_totals["footer_total"] == expected["total_amount"], f"{wb_totals['footer_total']} vs {expected['total_amount']}")
        check(checks, "reimbursement workbook summary total", wb_totals["summary_total"] == expected["total_amount"], f"{wb_totals['summary_total']} vs {expected['total_amount']}")
        check(checks, "reimbursement workbook summary count", wb_totals["summary_count"] == expected["expense_count"], f"{wb_totals['summary_count']} vs {expected['expense_count']}")
        summary_mismatches = summary_row_mismatches(wb_totals["summary_categories"], expected)
        check(checks, "reimbursement workbook category summaries", not summary_mismatches, "; ".join(summary_mismatches[:10]))
        check(
            checks,
            "reimbursement workbook category sum reconciles to total",
            sum(wb_totals["category_totals"].values(), Decimal("0.00")) == expected["total_amount"],
            f"{sum(wb_totals['category_totals'].values(), Decimal('0.00'))} vs {expected['total_amount']}",
        )
        check(checks, "reimbursement workbook image count", wb_totals["image_count"] == expected["screenshot_count"], f"{wb_totals['image_count']} vs {expected['screenshot_count']}")
        expected_image_rows = [row["sheet_row"] for row in wb_totals["rows"]]
        check(checks, "reimbursement workbook has one image on every detail row", wb_totals["image_rows"] == expected_image_rows, f"{wb_totals['image_rows']} vs {expected_image_rows}")

    finance_manifest_path = pack / "财务提交文件夹" / "manifest.json"
    check(checks, "finance folder manifest exists", finance_manifest_path.exists(), str(finance_manifest_path))
    if finance_manifest_path.exists():
        finance_manifest = read_json(finance_manifest_path)
        finance_integrity = manifest_integrity_issues(finance_manifest, finance_manifest_path)
        check(checks, "finance manifest amounts and source references are valid", not finance_integrity, "; ".join(finance_integrity[:10]))
        finance_expected = manifest_totals(finance_manifest)
        check(checks, "finance manifest total matches source manifest", finance_expected["total_amount"] == expected["total_amount"], f"{finance_expected['total_amount']} vs {expected['total_amount']}")
        check(checks, "finance manifest category totals match source manifest", finance_expected["category_totals"] == expected["category_totals"], f"{finance_expected['category_totals']} vs {expected['category_totals']}")
        finance_row_mismatches = detail_row_mismatches(manifest_detail_rows(finance_manifest), expected_rows)
        check(checks, "finance manifest rows match source manifest exactly", not finance_row_mismatches, "; ".join(finance_row_mismatches[:10]))
        if coverage_required:
            check(checks, "finance manifest invoice coverage matches source manifest", finance_manifest.get("invoice_coverage") == coverage, "invoice_coverage")

    expense_form = find_generated(summary, "费用报销单")
    form = manifest.get("reimbursement_form") or {}
    form_configured = bool(form.get("template"))
    check(checks, "expense form generated when configured", (not form_configured) or bool(expense_form and expense_form.exists()), str(expense_form or ""))
    check(checks, "expense form cell mapping configured", (not form_configured) or bool(form.get("cells")), str(form.get("cells") or ""))
    if expense_form and expense_form.exists() and expense_form.suffix.lower() in {".xlsx", ".xlsm"}:
        leftovers = form_placeholders(expense_form)
        check(checks, "expense form has no template placeholders", not leftovers, "; ".join(leftovers[:5]))
        if form.get("cells"):
            combined_for_form = find_generated(summary, "全部打印_一键打印")
            total_print_pages = pdf_page_count(combined_for_form) if combined_for_form and combined_for_form.exists() else None
            mismatches = form_mapping_mismatches(expense_form, form, manifest_context(manifest, expected, total_print_pages), manifest)
            check(checks, "expense form mapped values match manifest", not mismatches, "; ".join(mismatches[:5]))
        formula_cache_count, formula_cache_mismatches = form_formula_cache_mismatches(expense_form, expected["total_amount"])
        check(
            checks,
            "expense form formula caches match calculated values",
            formula_cache_count == 0 or not formula_cache_mismatches,
            f"checked={formula_cache_count}; " + "; ".join(formula_cache_mismatches[:5]),
        )

    screenshot_pages_actual = 0
    invoice_pages_actual = 0
    expense_form_pages_actual = 0

    expense_form_pdf = find_generated(summary, "费用报销单_打印")
    if form_configured:
        check(checks, "expense form print PDF exists", bool(expense_form_pdf and expense_form_pdf.exists()), str(expense_form_pdf or ""))
        if expense_form_pdf and expense_form_pdf.exists():
            expense_form_pages_actual = pdf_page_count(expense_form_pdf)
            check(checks, "expense form print PDF has pages", expense_form_pages_actual > 0, str(expense_form_pages_actual))
            form_pdf_text = "".join(page.extract_text() or "" for page in PdfReader(str(expense_form_pdf)).pages)
            normalized_form_pdf_text = re.sub(r"\s+", "", form_pdf_text)
            required_form_text = ["费用报销单", str(manifest.get("reimburser", ""))]
            missing_form_text = [value for value in required_form_text if value and value not in normalized_form_pdf_text]
            check(checks, "expense form print PDF contains required Chinese text", not missing_form_text, ", ".join(missing_form_text))

    screenshot_pdf = find_generated(summary, "费用截图_三张一页")
    check(checks, "screenshot print PDF exists", bool(screenshot_pdf and screenshot_pdf.exists()), str(screenshot_pdf or ""))
    if screenshot_pdf and screenshot_pdf.exists():
        expected_pages = math.ceil(expected["screenshot_count"] / 3) if expected["screenshot_count"] else 0
        screenshot_pages_actual = pdf_page_count(screenshot_pdf)
        check(checks, "screenshot print PDF pages", screenshot_pages_actual == expected_pages, f"{screenshot_pages_actual} vs {expected_pages}")

    invoice_pdf = find_generated(summary, "发票_一张一页")
    printable_invoice_documents = expected["invoice_count"] + expected["supporting_document_count"]
    if printable_invoice_documents:
        check(checks, "invoice print PDF exists", bool(invoice_pdf and invoice_pdf.exists()), str(invoice_pdf or ""))
        if invoice_pdf and invoice_pdf.exists():
            invoice_pages_actual = pdf_page_count(invoice_pdf)
            check(checks, "invoice print PDF pages", invoice_pages_actual >= printable_invoice_documents, f"{invoice_pages_actual} >= {printable_invoice_documents}")
    else:
        check(checks, "invoice print PDF omitted when no invoices", not invoice_pdf, str(invoice_pdf or ""))

    combined_print_pdf = find_generated(summary, "全部打印_一键打印")
    check(checks, "combined one-click print PDF exists", bool(combined_print_pdf and combined_print_pdf.exists()), str(combined_print_pdf or ""))
    if combined_print_pdf and combined_print_pdf.exists():
        combined_reader = PdfReader(str(combined_print_pdf))
        combined_pages = len(combined_reader.pages)
        expected_combined_pages = expense_form_pages_actual + screenshot_pages_actual + invoice_pages_actual
        check(checks, "combined one-click print PDF pages", combined_pages == expected_combined_pages, f"{combined_pages} vs {expected_combined_pages}")
        page_number_issues = []
        for page_index, page in enumerate(combined_reader.pages, start=1):
            text = re.sub(r"\s+", " ", page.extract_text() or "")
            if f"{page_index} / {combined_pages}" not in text:
                page_number_issues.append(str(page_index))
        check(checks, "combined one-click print PDF has continuous page numbers", not page_number_issues, ", ".join(page_number_issues[:10]))
        if expense_form_pdf and expense_form_pdf.exists():
            form_first_page = PdfReader(str(expense_form_pdf)).pages[0]
            combined_first_page = combined_reader.pages[0]
            form_text = re.sub(r"\s+", "", form_first_page.extract_text() or "")
            combined_text = re.sub(r"\s+", "", combined_first_page.extract_text() or "")
            check(checks, "combined one-click print PDF starts with expense form", bool(form_text and form_text in combined_text), "page 1")

    finance_zip = Path(summary.get("finance_zip", ""))
    check(checks, "finance zip exists", finance_zip.exists(), str(finance_zip))
    if finance_zip.exists():
        report = zip_report(finance_zip)
        check(checks, "finance zip has standard folders", all([report["has_tables_dir"], report["has_screenshots_dir"], report["has_invoices_dir"], report["has_print_dir"]]), str({k: report[k] for k in ["has_tables_dir", "has_screenshots_dir", "has_invoices_dir", "has_print_dir"]}))
        check(checks, "finance zip is clean", not report["bad"], "; ".join(report["bad"][:5]))
        check(checks, "finance zip screenshot files", report["screenshot_files"] == expected["screenshot_count"], f"{report['screenshot_files']} vs {expected['screenshot_count']}")
        if manifest.get("workflow_version", 1) >= 2:
            screenshot_ids_ok = all(re.match(r"FY\d{3,}_", Path(name).name) for name in report["screenshot_names"])
            check(checks, "finance zip screenshot names use voucher IDs", screenshot_ids_ok, ", ".join(Path(name).name for name in report["screenshot_names"][:5]))
        expected_finance_invoice_files = expected["invoice_count"] + expected["supporting_document_count"]
        check(checks, "finance zip invoice and supporting files", report["invoice_files"] == expected_finance_invoice_files, f"{report['invoice_files']} vs {expected_finance_invoice_files}")
        check(checks, "finance zip supporting document files", report["supporting_document_files"] == expected["supporting_document_count"], f"{report['supporting_document_files']} vs {expected['supporting_document_count']}")
        check(checks, "finance zip table files", report["table_files"] >= 1, str(report["table_files"]))
        check(checks, "finance zip print files", report["print_files"] >= 2, str(report["print_files"]))
        check(checks, "finance zip has delivery checklist", report["has_delivery_checklist"], "交付清单.md")
        if manifest.get("workflow_version", 1) >= 2:
            check(checks, "finance zip has expense confirmation report", report["has_expense_confirmation"], "报销金额确认单.md")
            check(checks, "finance zip has expense exception report", report["has_expense_exceptions"], "报销异常清单.md")
            check(checks, "finance zip has DingTalk gap report", report["has_approval_gap"], "钉钉信息缺口.md")
        check(checks, "finance zip has audit summary", report["has_audit_summary"], "报销审计摘要.md")
        if coverage_required:
            check(checks, "finance zip has invoice coverage JSON", report["has_invoice_coverage_report"], "invoice_coverage.json")
            check(checks, "finance zip has invoice gap report", report["has_invoice_gap_report"], "发票缺口清单.md")
            check(
                checks,
                "finance zip replacement invoice count",
                report["replacement_invoice_files"] == replacement_invoice_count,
                f"{report['replacement_invoice_files']} vs {replacement_invoice_count}",
            )

    failed = [item for item in checks if item["status"] != "passed"]
    reconciliation = {
        "currency": "CNY",
        "precision": "0.01",
        "expense_count": expected["expense_count"],
        "manifest_total": f"{expected['total_amount']:.2f}",
        "pack_summary_total": f"{summary_total:.2f}",
        "workbook_detail_total": f"{wb_totals['total_amount']:.2f}" if wb_totals else "",
        "workbook_summary_total": f"{wb_totals['summary_total']:.2f}" if wb_totals and wb_totals["summary_total"] is not None else "",
        "workbook_footer_total": f"{wb_totals['footer_total']:.2f}" if wb_totals and wb_totals["footer_total"] is not None else "",
        "category_totals": {category: f"{amount:.2f}" for category, amount in expected["category_totals"].items()},
    }
    invoice_reconciliation = {
        "policy": invoice_policy,
        "status": (coverage.get("status") or "not_analyzed") if coverage_required else "not_required",
        "required_amount": f"{coverage_amounts.get('required_amount', expected['total_amount']):.2f}" if coverage_required else "0.00",
        "recognized_amount": f"{coverage_amounts.get('recognized_amount', Decimal('0.00')):.2f}" if coverage_required else "0.00",
        "missing_amount": f"{coverage_amounts.get('missing_amount', expected['total_amount']):.2f}" if coverage_required else "0.00",
        "replacement_invoice_count": replacement_invoice_count,
        "replacement_invoice_amount": f"{coverage_amounts.get('replacement_invoice_amount', Decimal('0.00')):.2f}" if coverage_required else "0.00",
    }
    result = {
        "status": "failed" if failed else "passed",
        "pack": str(pack),
        "summary": str(summary_path),
        "manifest": str(manifest_path or ""),
        "financial_reconciliation": reconciliation,
        "invoice_reconciliation": invoice_reconciliation,
        "checks": checks,
    }
    output = Path(args.out).expanduser().resolve() if args.out else pack / "verification.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if failed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
